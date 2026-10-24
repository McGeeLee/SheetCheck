"""Bounded, in-memory CSV/XLSX import with original types and source positions.

No workbook formula is evaluated and no workbook relationship is followed over
network. XLSX XML is preflighted before openpyxl sees it. Resource limits are
acceptance limits, not a claim that all inputs fit a particular memory budget.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import posixpath
import re
import warnings
import zipfile
from dataclasses import replace
from pathlib import PurePosixPath
from typing import BinaryIO

from defusedxml.ElementTree import iterparse
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries

from .errors import InputError
from .models import Cell, FileSpec, Region, Side, SourcePosition, SourceRow, SourceTable

_SHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
_MAX_XML_DEPTH = 128
_CSV_MAX_FIELD_CHARS = 128 * 1024
_EXCEL_MAX_ROW = 1_048_576
_EXCEL_MAX_COL = 16_384
_CELL_COORDINATE = re.compile(r"[A-Za-z]{1,3}[1-9][0-9]{0,6}\Z")
_FORBIDDEN_PARTS = (
    "vbaproject", "vba", "activex", "embeddings", "externallinks", "macrosheets",
    "dialogsheets", "customui", "connections", "querytables", "webextensions",
    "ctrlprops", "oleobject", "attachedtemplate", "customxml",
)
_FORBIDDEN_TYPES = (
    "macroenabled", "vbaproject", "activex", "oleobject", "externalLink".lower(),
    "macrosheet", "dialogsheet", "customui", "connections", "querytable", "webextension",
    "control", "attachedtemplate", "customxml",
)


def _fail(code: str, message: str) -> None:
    raise InputError(code, message)


def _source_id(side: Side, digest: str, sheet: str | None, position: SourcePosition) -> str:
    identity = [str(side), digest, sheet if sheet is not None else "CSV", position.row,
                position.physical_start, position.physical_end]
    return hashlib.sha256(json.dumps(identity, ensure_ascii=False, separators=(",", ":")).encode("utf-8")).hexdigest()


def _check_spec(data: bytes, spec: FileSpec) -> str:
    if not isinstance(data, bytes):
        _fail("INPUT_BYTES_REQUIRED", "输入必须为文件字节；核心不读取文件路径。")
    if not isinstance(spec, FileSpec):
        _fail("FILE_SPEC_INVALID", "请提供有效的文件参数。")
    try:
        Side(spec.side)
    except (TypeError, ValueError):
        _fail("FILE_SIDE_INVALID", "文件角色必须为订单或收款。")
    if not isinstance(spec.filename, str) or not spec.filename:
        _fail("FILE_TYPE_UNSUPPORTED", "请选择 CSV 或 XLSX 文件。")
    extension = PurePosixPath(spec.filename).suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        _fail("FILE_TYPE_UNSUPPORTED", "仅支持 CSV 和 XLSX，不支持 XLS、XLSM 或压缩包。")
    limits = spec.limits
    for name in ("max_file_bytes", "max_records", "max_columns", "max_uncompressed_bytes", "max_zip_entries", "max_compression_ratio"):
        value = getattr(limits, name, None)
        if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
            _fail("LIMITS_INVALID", "容量限制必须为正整数。")
    if len(data) > limits.max_file_bytes:
        _fail("FILE_SIZE_LIMIT", "文件超过允许大小；不会截断后继续。")
    if not data:
        _fail("FILE_EMPTY", "文件为空。")
    if isinstance(spec.header_row, bool) or not isinstance(spec.header_row, int) or spec.header_row < 1:
        _fail("HEADER_ROW_INVALID", "表头行必须为从 1 开始的整数。")
    if spec.header_row - 1 > limits.max_records or spec.header_row > _EXCEL_MAX_ROW:
        _fail("PREAMBLE_LIMIT", "表头之前的前言行数量超过记录上限。")
    return extension


def read_source(data: bytes, file_spec: FileSpec) -> SourceTable:
    """Read exactly one selected table; any file-level error aborts the import.

    CSV ``header_row`` counts logical CSV records, not physical text lines.
    A single-sheet XLSX may omit ``sheet``; a multi-sheet XLSX must select it.
    Blank CSV lines are retained as all-empty records with the header width.
    """
    extension = _check_spec(data, file_spec)
    digest = hashlib.sha256(data).hexdigest()
    if extension == ".csv":
        return _read_csv(data, file_spec, digest)
    return _read_xlsx(data, file_spec, digest)


def list_worksheets(data: bytes, file_spec: FileSpec) -> tuple[str, ...]:
    """Discover ordinary XLSX worksheets after the same safe import preflight.

    This metadata-only operation does not load cells, evaluate formulas, follow
    external links, or skip inspection of unselected ZIP/XML members. Hidden
    worksheets remain selectable. CSV has no workbook sheets.
    """
    if _check_spec(data, file_spec) != ".xlsx":
        _fail("XLSX_REQUIRED", "工作表选择仅适用于 XLSX 文件。")
    if not data.startswith(b"PK\x03\x04"):
        _fail("FILE_SIGNATURE_INVALID", "文件签名不是未加密 XLSX；不支持加密或损坏文件。")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            _preflight_zip(archive, file_spec)
            content_types = _small_xml_items(archive, "[Content_Types].xml", "{http://schemas.openxmlformats.org/package/2006/content-types}Override")
            workbook_types = [item.get("ContentType") for item in content_types if item.get("PartName") == "/xl/workbook.xml"]
            if workbook_types != [_XLSX_CONTENT_TYPE]:
                _fail("XLSX_CONTENT_TYPE", "文件不是普通 XLSX 工作簿。")
            sheets = _small_xml_items(archive, "xl/workbook.xml", f"{{{_SHEET_NS}}}sheet")
            relations = _small_xml_items(archive, "xl/_rels/workbook.xml.rels", f"{{{_REL_NS}}}Relationship")
            names: list[str] = []
            all_names = [sheet.get("name") for sheet in sheets]
            if any(not name for name in all_names) or len(set(all_names)) != len(all_names):
                _fail("XLSX_CONTAINER_INVALID", "工作表名称缺失或不唯一，请重新导出普通工作簿。")
            for sheet in sheets:
                relation_id = sheet.get(f"{{{_DOC_REL_NS}}}id")
                matches = [item for item in relations if item.get("Id") == relation_id]
                if len(matches) != 1:
                    _fail("XLSX_CONTAINER_INVALID", "工作表部件引用无效。")
                if not matches[0].get("Type", "").endswith("/worksheet"):
                    continue
                # Reuse the selected-table relationship and path validation.
                selected_name, _ = _worksheet_path(archive, sheet, relations)
                names.append(selected_name)
            if not names:
                _fail("XLSX_NO_SHEETS", "XLSX 不含可读取的普通工作表。")
            return tuple(names)
    except InputError:
        raise
    except Exception:
        raise InputError("XLSX_INVALID", "XLSX 损坏、含不安全 XML 或结构不受支持；请重新导出纯值副本。") from None


def _strict_quotes(text: str, delimiter: str) -> None:
    # csv.reader(strict=True) still accepts bare quotes inside an unquoted
    # field. Require conventional quoting instead of silently accepting them.
    state = "start"
    for char in text:
        if state == "quoted":
            if char == '"':
                state = "closed"
        elif state == "closed":
            if char == '"':
                state = "quoted"
            elif char == delimiter or char in "\r\n":
                state = "start"
            else:
                _fail("CSV_PARSE_ERROR", "CSV 引号格式不合法；请重新导出标准 CSV。")
        elif char == '"':
            if state != "start":
                _fail("CSV_PARSE_ERROR", "CSV 未加引号的字段中含裸引号。")
            state = "quoted"
        elif char == delimiter or char in "\r\n":
            state = "start"
        else:
            state = "plain"
    if state == "quoted":
        _fail("CSV_PARSE_ERROR", "CSV 引号未闭合。")


def _read_csv(data: bytes, spec: FileSpec, digest: str) -> SourceTable:
    if spec.sheet is not None:
        _fail("CSV_SHEET_INVALID", "CSV 不包含工作表；请清除工作表选择。")
    if spec.encoding not in {"utf-8", "utf-8-sig", "gb18030"}:
        _fail("CSV_ENCODING_UNSUPPORTED", "CSV 仅支持 UTF-8、UTF-8 BOM 或明确选择的 GB18030。")
    if spec.delimiter not in {",", "\t", ";"}:
        _fail("CSV_DELIMITER_UNSUPPORTED", "CSV 分隔符仅支持逗号、制表符或分号。")
    if (data.startswith((b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"))
            or zipfile.is_zipfile(io.BytesIO(data))):
        _fail("FILE_SIGNATURE_INVALID", "文件签名与 CSV 类型不符。")
    try:
        text = data.decode("utf-8-sig" if spec.encoding == "utf-8" else spec.encoding, errors="strict")
    except UnicodeDecodeError:
        _fail("CSV_DECODE_ERROR", "CSV 解码失败；请确认 UTF-8 或明确选择 GB18030。")
    _strict_quotes(text, spec.delimiter)
    records: list[SourceRow] = []
    headers: tuple[Cell, ...] | None = None
    previous_end = 0
    # Avoid global csv.field_size_limit mutation: stdlib's 128 KiB field limit
    # is an additional documented rejection boundary, never silent truncation.
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=spec.delimiter, strict=True)
    try:
        for logical_number, values in enumerate(reader, start=1):
            physical_start, physical_end = previous_end + 1, reader.line_num
            previous_end = physical_end
            if any(len(value) > _CSV_MAX_FIELD_CHARS for value in values):
                _fail("CSV_FIELD_LIMIT", "CSV 单字段超过 128 Ki 字符限制。")
            if len(values) > spec.limits.max_columns:
                _fail("COLUMN_LIMIT", "列数超过允许上限。")
            if logical_number < spec.header_row:
                continue
            if logical_number == spec.header_row:
                headers = tuple(Cell(value, "text") for value in (values or [""]))
                continue
            if headers is None:  # defensive; loop order guarantees the header
                _fail("HEADER_ROW_MISSING", "所选表头行不存在。")
            if len(records) >= spec.limits.max_records:
                _fail("RECORD_LIMIT", "数据记录超过允许上限；不会截断后继续。")
            if not values:
                values = [""] * len(headers)
            elif len(values) != len(headers):
                _fail("CSV_COLUMN_COUNT", "CSV 数据记录与表头列数不一致。")
            position = SourcePosition(logical_number, physical_start=physical_start, physical_end=physical_end)
            records.append(SourceRow(_source_id(spec.side, digest, None, position), spec.side, position,
                                     tuple(Cell(value, "text") for value in values)))
    except csv.Error:
        _fail("CSV_PARSE_ERROR", "CSV 解析失败或单字段超过解析器容量；请整理源文件。")
    if headers is None:
        _fail("HEADER_ROW_MISSING", "所选表头行不存在。")
    return SourceTable(spec, digest, headers, tuple(records), spec.header_row - 1)


def _xml_events(stream: BinaryIO):
    return iterparse(stream, events=("start", "end"), forbid_dtd=True, forbid_entities=True, forbid_external=True)


def _preflight_xml(stream: BinaryIO, *, relationships: bool = False, content_types: bool = False) -> None:
    stack = []
    for event, element in _xml_events(stream):
        if event == "start":
            stack.append(element)
            if len(stack) > _MAX_XML_DEPTH:
                _fail("XLSX_XML_DEPTH", "XLSX XML 嵌套层级过深。")
            if relationships and element.tag == f"{{{_REL_NS}}}Relationship":
                target_mode = element.get("TargetMode", "").lower()
                target = element.get("Target", "")
                relation_type = element.get("Type", "").lower()
                if target_mode == "external" or re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", target) or target.startswith("//"):
                    _fail("XLSX_EXTERNAL_LINK", "XLSX 含外部链接；请从可信来源导出无外链的纯值副本。")
                if any(marker in relation_type for marker in _FORBIDDEN_TYPES):
                    _fail("XLSX_ACTIVE_CONTENT", "XLSX 含不支持的活动内容或数据连接。")
            if content_types:
                kind = element.get("ContentType", "").lower()
                if any(marker in kind for marker in _FORBIDDEN_TYPES):
                    _fail("XLSX_ACTIVE_CONTENT", "XLSX 含宏、活动内容或外链部件。")
        else:
            stack.pop()
            if stack:
                stack[-1].remove(element)
            element.clear()


def _preflight_zip(archive: zipfile.ZipFile, spec: FileSpec) -> None:
    infos = archive.infolist()
    limits = spec.limits
    if len(infos) > limits.max_zip_entries:
        _fail("XLSX_ZIP_ENTRY_LIMIT", "XLSX 压缩包条目数超过允许上限。")
    names: set[str] = set()
    total = 0
    for info in infos:
        path = info.filename
        lower = path.lower()
        segments = PurePosixPath(path).parts
        if (path.startswith("/") or "\\" in path or any(part in {"..", "."} for part in path.split("/"))
                or re.match(r"^[a-zA-Z]:", path) or "\x00" in path or lower in names):
            _fail("XLSX_CONTAINER_INVALID", "XLSX 容器路径不合法或包含重复部件。")
        names.add(lower)
        if info.flag_bits & 1:
            _fail("XLSX_ENCRYPTED", "不支持加密 XLSX。")
        if info.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED}:
            _fail("XLSX_COMPRESSION_UNSUPPORTED", "XLSX 使用不支持的压缩方式。")
        total += info.file_size
        if total > limits.max_uncompressed_bytes:
            _fail("XLSX_UNCOMPRESSED_LIMIT", "XLSX 累计解压大小超过允许上限。")
        if info.file_size > max(info.compress_size, 1) * limits.max_compression_ratio:
            _fail("XLSX_COMPRESSION_RATIO", "XLSX 条目压缩比超过允许上限。")
        if any(marker in lower for marker in _FORBIDDEN_PARTS) or lower.endswith((".bin", ".exe", ".dll", ".js", ".vbs")):
            _fail("XLSX_ACTIVE_CONTENT", "XLSX 含不支持的嵌入、活动或外链部件。")
        if not segments:
            _fail("XLSX_CONTAINER_INVALID", "XLSX 容器中存在无效部件。")
    required = {"[content_types].xml", "_rels/.rels", "xl/workbook.xml", "xl/_rels/workbook.xml.rels"}
    if not required <= names:
        _fail("XLSX_CONTAINER_INVALID", "文件不是受支持的完整 XLSX 容器。")
    # Reading every entry checks CRC and actual deflate integrity. Stream XML
    # and remove processed elements, so unselected sheets do not form trees.
    for info in infos:
        if info.is_dir():
            continue
        lower = info.filename.lower()
        with archive.open(info) as stream:
            if lower.endswith((".xml", ".rels")):
                _preflight_xml(stream, relationships=lower.endswith(".rels"), content_types=lower == "[content_types].xml")
            else:
                while stream.read(1024 * 1024):
                    pass


def _small_xml_items(archive: zipfile.ZipFile, name: str, tag: str) -> list[dict[str, str]]:
    result = []
    with archive.open(name) as stream:
        stack = []
        for event, element in _xml_events(stream):
            if event == "start":
                stack.append(element)
                if element.tag == tag:
                    result.append(dict(element.attrib))
            else:
                stack.pop()
                if stack:
                    stack[-1].remove(element)
                element.clear()
    return result


def _selected_sheet(archive: zipfile.ZipFile, spec: FileSpec) -> tuple[str, str]:
    sheets = _small_xml_items(archive, "xl/workbook.xml", f"{{{_SHEET_NS}}}sheet")
    if not sheets:
        _fail("XLSX_NO_SHEETS", "XLSX 不含可读取工作表。")
    if spec.sheet is None:
        if len(sheets) != 1:
            _fail("SHEET_REQUIRED", "文件含多个工作表，请明确选择一个。")
        selected = sheets[0]
    else:
        matches = [sheet for sheet in sheets if sheet.get("name") == spec.sheet]
        if len(matches) != 1:
            _fail("SHEET_NOT_FOUND", "所选工作表不存在或名称不唯一。")
        selected = matches[0]
    relations = _small_xml_items(archive, "xl/_rels/workbook.xml.rels", f"{{{_REL_NS}}}Relationship")
    return _worksheet_path(archive, selected, relations)


def _worksheet_path(archive: zipfile.ZipFile, selected: dict[str, str], relations: list[dict[str, str]]) -> tuple[str, str]:
    relation_id = selected.get(f"{{{_DOC_REL_NS}}}id")
    matches = [item for item in relations if item.get("Id") == relation_id]
    if len(matches) != 1 or not matches[0].get("Type", "").endswith("/worksheet"):
        _fail("SHEET_TYPE_UNSUPPORTED", "所选对象不是普通工作表。")
    target = matches[0].get("Target", "")
    path = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join("xl", target))
    if not path.startswith("xl/") or path not in archive.namelist():
        _fail("XLSX_CONTAINER_INVALID", "工作表部件引用无效。")
    if not selected.get("name"):
        _fail("XLSX_CONTAINER_INVALID", "工作表名称无效。")
    return selected["name"], path


def _region(reference: str) -> Region:
    try:
        min_col, min_row, max_col, max_row = range_boundaries(reference)
        if (any(value is None for value in (min_col, min_row, max_col, max_row))
                or not 1 <= min_row <= max_row <= _EXCEL_MAX_ROW
                or not 1 <= min_col <= max_col <= _EXCEL_MAX_COL):
            raise ValueError
        return Region(min_row, max_row, min_col, max_col)
    except (TypeError, ValueError):
        _fail("XLSX_RANGE_INVALID", "XLSX 中的公式或合并区域引用不合法。")


def _scan_sheet(archive: zipfile.ZipFile, path: str, spec: FileSpec):
    max_row = max_col = 0
    merged: list[Region] = []
    formulas: list[Region] = []
    hidden: set[int] = set()
    current_row = previous_row = last_col = 0
    stack = []
    with archive.open(path) as stream:
        for event, element in _xml_events(stream):
            local = element.tag.rsplit("}", 1)[-1]
            if event == "start":
                stack.append(element)
                if local == "row":
                    try:
                        current_row = int(element.get("r", "0"))
                    except ValueError:
                        _fail("XLSX_POSITION_INVALID", "工作表行位置不合法。")
                    if not previous_row < current_row <= _EXCEL_MAX_ROW:
                        _fail("XLSX_POSITION_INVALID", "工作表行位置重复、乱序或越界。")
                    previous_row, last_col = current_row, 0
                    if element.get("hidden", "").lower() in {"1", "true"}:
                        hidden.add(current_row)
                continue
            if local == "c":
                coordinate = element.get("r", "")
                if not _CELL_COORDINATE.fullmatch(coordinate):
                    _fail("XLSX_POSITION_INVALID", "工作表单元格位置不合法。")
                row, col = coordinate_to_tuple(coordinate)
                if row != current_row or not last_col < col <= _EXCEL_MAX_COL:
                    _fail("XLSX_POSITION_INVALID", "工作表单元格重复、乱序或越界。")
                last_col = col
                formula = element.find(f"{{{_SHEET_NS}}}f")
                value = element.find(f"{{{_SHEET_NS}}}v")
                inline = element.find(f"{{{_SHEET_NS}}}is")
                has_value = value is not None and value.text is not None
                has_inline = inline is not None and any(item.text is not None for item in inline.iter(f"{{{_SHEET_NS}}}t"))
                if formula is not None:
                    region = _region(formula.get("ref") or coordinate)
                    formulas.append(region)
                    max_row, max_col = max(max_row, row, region.max_row), max(max_col, col, region.max_col)
                elif has_value or has_inline:
                    max_row, max_col = max(max_row, row), max(max_col, col)
                if max_col > spec.limits.max_columns:
                    _fail("COLUMN_LIMIT", "工作表实际内容列数超过允许上限。")
                if max_row - spec.header_row > spec.limits.max_records:
                    _fail("RECORD_LIMIT", "表头后的实际数据跨度超过记录上限，包含中间空行。")
            elif local == "mergeCell":
                merged.append(_region(element.get("ref", "")))
            stack.pop()
            # Preserve descendants only until their enclosing cell is examined.
            if not any(parent.tag == f"{{{_SHEET_NS}}}c" for parent in stack):
                if stack:
                    stack[-1].remove(element)
                element.clear()
    if max_row < spec.header_row or max_col == 0:
        _fail("HEADER_ROW_MISSING", "所选表头行不在工作表实际内容范围内。")
    hidden = {row for row in hidden if spec.header_row < row <= max_row}
    return max_row, max_col, tuple(merged), tuple(formulas), hidden


def _cell(raw) -> Cell:
    value = raw.value
    if raw.data_type == "f":
        # Special formula objects have implementation-specific reprs; retain a
        # stable marker for their raw value and collect ranges in the XML scan.
        return Cell(value if isinstance(value, str) else "[FORMULA_REGION]", "formula")
    if value is None:
        return Cell(None, "empty")
    if raw.data_type == "e":
        return Cell(value, "error")
    if raw.data_type == "b" or isinstance(value, bool):
        return Cell(value, "boolean")
    if raw.is_date or raw.data_type == "d":
        return Cell(value, "date")
    if raw.data_type == "n":
        return Cell(value, "number")
    if isinstance(value, str):
        return Cell(value, "text")
    _fail("XLSX_CELL_TYPE", "XLSX 包含不支持的单元格类型。")


def _read_xlsx(data: bytes, spec: FileSpec, digest: str) -> SourceTable:
    if not data.startswith(b"PK\x03\x04"):
        _fail("FILE_SIGNATURE_INVALID", "文件签名不是未加密 XLSX；不支持加密或损坏文件。")
    workbook = None
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            _preflight_zip(archive, spec)
            content_types = _small_xml_items(archive, "[Content_Types].xml", "{http://schemas.openxmlformats.org/package/2006/content-types}Override")
            workbook_types = [item.get("ContentType") for item in content_types if item.get("PartName") == "/xl/workbook.xml"]
            if workbook_types != [_XLSX_CONTENT_TYPE]:
                _fail("XLSX_CONTENT_TYPE", "文件不是普通 XLSX 工作簿。")
            name, path = _selected_sheet(archive, spec)
            max_row, max_col, merged, formulas, hidden = _scan_sheet(archive, path, spec)
        # Some openpyxl warnings contain raw numeric values (for example an
        # invalid date serial). Suppress those library warnings so source data
        # cannot leak through console logs; affected cell types stay invalid.
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", category=UserWarning, module=r"openpyxl(?:\..*)?")
            # read_only avoids materialising formatted-but-empty sheet ranges.
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            sheet = workbook[name]
            sheet.reset_dimensions()
            effective_spec = replace(spec, sheet=name)
            rows = []
            headers = None
            for number, raw_cells in enumerate(sheet.iter_rows(min_row=spec.header_row, max_row=max_row, min_col=1, max_col=max_col), spec.header_row):
                cells = tuple(_cell(raw) for raw in raw_cells)
                if number == spec.header_row:
                    headers = cells
                    continue
                position = SourcePosition(number, sheet=name, hidden=number in hidden)
                rows.append(SourceRow(_source_id(spec.side, digest, name, position), spec.side, position, cells))
            if headers is None:
                _fail("HEADER_ROW_MISSING", "所选表头行不存在。")
            return SourceTable(effective_spec, digest, headers, tuple(rows), spec.header_row - 1, merged, formulas, len(hidden))
    except InputError:
        raise
    except Exception:
        # Do not include parser messages: they may contain source values/paths.
        raise InputError("XLSX_INVALID", "XLSX 损坏、含不安全 XML 或结构不受支持；请重新导出纯值副本。") from None
    finally:
        if workbook is not None:
            workbook.close()

"""Static XLSX report with reversible text encoding and mandatory readback audit.

No formulas, links, user paths, or disk-backed temporary worksheet files.
The in-memory writer adapter is deliberately pinned/tested against openpyxl 3.1.5.
"""
from __future__ import annotations

from dataclasses import asdict
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from itertools import zip_longest
import gc
import json
import posixpath
import re
import unicodedata
from zipfile import ZipFile, ZIP_DEFLATED

from defusedxml.ElementTree import iterparse
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.spreadsheet_drawing import SpreadsheetDrawing
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet._writer import WorksheetWriter
from openpyxl.writer.excel import ExcelWriter

from .audit import assert_invariants
from .errors import ExportError, IntegrityError, PROBLEM_HELP
from .models import ReconciliationResult, Side, Status
from .report_schema import (BUSINESS_COLUMNS, ISSUE_COLUMNS, LEDGER_COLUMNS, REPORT_VERSION,
                            RUN_HEADERS, SHEET_NAMES, TEXT_ESCAPE_PREFIX)

MAX_REPORT_BYTES = 128 * 1024 * 1024
MAX_REPORT_XML_BYTES = 512 * 1024 * 1024
MAX_CELL_UTF16 = 32767
_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_COLUMN_SETS = {"业务结果": BUSINESS_COLUMNS, "异常明细": BUSINESS_COLUMNS,
                "数据问题": ISSUE_COLUMNS, "源行账本": LEDGER_COLUMNS}
_HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
_BODY_FONT = Font(name="Arial", size=11, color="182338")
_HEADER_FILL = PatternFill("solid", fgColor="243B63")
_TEXT_ALIGN = Alignment(vertical="center", wrap_text=True)
_NUMBER_ALIGN = Alignment(horizontal="right", vertical="center")


def _text_length(value: str) -> None:
    if len(value.encode("utf-16-le", errors="surrogatepass")) // 2 > MAX_CELL_UTF16:
        raise ExportError("EXPORT_TEXT_LIMIT", "转义后的字段超过 Excel 单元格 32767 个 UTF-16 单位；已中止，未截断。请缩短源字段后重试。")


def encode_text(value: str) -> str:
    """Encode only unsafe/ambiguous user text; never mutate the source string."""
    invalid_xml = any((ord(c) < 32 and c not in "\t\n") or c == "\r"
                      or unicodedata.category(c) == "Cs" or (ord(c) & 0xFFFF) in {0xFFFE, 0xFFFF}
                      for c in value)
    dangerous = bool(value) and (value.lstrip().startswith(("=", "+", "-", "@"))
                                 or unicodedata.category(value[0]) in {"Cc", "Cf", "Cs"})
    encoded = (TEXT_ESCAPE_PREFIX + json.dumps(value, ensure_ascii=True)
               if not value or value.startswith(TEXT_ESCAPE_PREFIX) or invalid_xml or dangerous else value)
    _text_length(encoded)
    return encoded


def decode_text(value: str) -> str:
    if not value.startswith(TEXT_ESCAPE_PREFIX):
        return value
    try:
        decoded = json.loads(value[len(TEXT_ESCAPE_PREFIX):])
        if not isinstance(decoded, str):
            raise ValueError
        return decoded
    except (ValueError, TypeError):
        raise ExportError("EXPORT_ESCAPE_INVALID", "文本转义格式无效。") from None


def _raw_text(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date, time)):
        return encode_text(value.isoformat())
    # Bool/date/error/numeric source values remain text, with separate raw type.
    return encode_text(str(value))


def _money(minor):
    if minor is None:
        return None
    sign = "-" if minor < 0 else ""
    return f"{sign}{abs(minor) // 100}.{abs(minor) % 100:02d}"


def _json(value):
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"), allow_nan=False)


def _outcome(result):
    counts = result.manifest.status_counts
    issues = counts[Status.DATA_ISSUE]
    exceptions = sum(counts[s] for s in (Status.AMOUNT_MISMATCH, Status.ONLY_ORDER, Status.ONLY_PAYMENT))
    if not result.records:
        return "无可核对业务记录，仍有数据问题" if issues else "无可核对业务记录"
    if exceptions:
        return "存在差额或缺项，仍有数据问题" if issues else "存在差额或缺项"
    return "可核对部分一致，仍有数据问题" if issues else "本次可核对部分全部一致"


def _run_rows(result):
    manifest = result.manifest
    yield ("概览", "outcome", _outcome(result), "仅描述本次有效输入的核对结论。")
    yield ("概览", "currency", manifest.rules.currency, "金额显示列保留两位小数，分值列为整数。")
    yield ("概览", "period", encode_text(manifest.rules.period), "用户声明的核对期间。")
    for key in ("app_version", "run_id", "fingerprint", "started_at", "finished_at"):
        yield ("运行", key, encode_text(str(getattr(manifest, key))), "时间为带时区的 ISO 8601；指纹标识输入及规则。" if key.endswith("_at") else None)
    yield ("运行", "report_version", REPORT_VERSION, "静态报告格式版本。")
    for key, value in asdict(manifest.rules).items():
        output = ("是" if value else "否") if isinstance(value, bool) else encode_text(value) if isinstance(value, str) else value
        yield ("规则", f"rules.{key}", output, "编号精确匹配；差额 = 收款减订单；不舍入。" if key == "tolerance_minor" else None)
    yield ("风险", "unidentifiable_id_risk", "是" if manifest.unidentifiable_id_risk else "否", "无法识别编号的问题行可能隐藏对方记录。")
    yield ("风险", "warnings", _json(manifest.warnings), "缺项仅指本次有效集合中未配对，需要复核期间与输入完整性。")
    for status, value in manifest.status_counts.items():
        yield ("数量", f"status_counts.{status}", value, "源行数" if status in (Status.DATA_ISSUE, Status.SKIPPED_EMPTY) else "业务结果条数；配对对应两条源行")
    yield ("数量", "business_result_count", len(result.records), "结果条数不能当作源行数。")
    yield ("数量", "ledger_count", len(result.ledger), "每条输入源记录恰好一条账本记录。")
    for side, summary in manifest.summaries.items():
        for key, value in asdict(summary).items():
            explanation = "整数分" if key.endswith("_minor") else "源行数"
            if key == "issue_parseable_minor":
                explanation = "仅问题行可解析部分（分）；另有金额无法解析时，问题金额不完整。"
            precise_value = str(value) if abs(value) >= 10**15 else value
            if isinstance(precise_value, str):
                explanation += "；超过 Excel 15 位数字精度，整数分以精确文本保存。"
            yield ("订单汇总" if side == Side.ORDER else "收款汇总", f"summaries.{side}.{key}", precise_value, explanation)
            if key.endswith("_minor"):
                yield ("金额显示", f"summaries.{side}.{key}.display", _money(value), "两位小数字符串；对应上行整数分。")
    for table in result.tables:
        side, source = str(table.source.spec.side), table.source
        group = "订单输入" if side == Side.ORDER else "收款输入"
        for key, value in manifest.inputs[list(Side).index(table.source.spec.side)].items():
            if key == "mapping":
                for k, v in value.items():
                    yield (group, f"inputs.{side}.mapping.{k}", v, "列号从 1 开始。")
            elif isinstance(value, dict):
                yield (group, f"inputs.{side}.{key}", _json(value), None)
            else:
                yield (group, f"inputs.{side}.{key}", encode_text(value) if isinstance(value, str) else value, None)
        for label, column in (("id", table.mapping.id_column), ("amount", table.mapping.amount_column)):
            header = source.headers[column - 1]
            yield (group, f"inputs.{side}.{label}_header", _raw_text(header.value), f"原表头类型：{header.kind}。映射以列号为准。")
    yield ("解释", "delta_direction", "收款金额减订单金额", "正负只代表本次配对数值大小，不直接认定应退、欠款或错误。")
    yield ("解释", "source_navigation", "按 source_id 关联源行账本；按 result_id 关联业务结果。", "同一身份在多个页签出现仅为引用，不能重复汇总。")
    yield ("安全", "text_encoding", "编码格式为 SC_ESC: 加 JSON 字符串", "若字段以 SC_ESC: 开始，仅解码一次其后 JSON 字符串。包括原本以此前缀开头的文本；不对内存原值做修改。")
    yield ("安全", "empty_values", "空单元格表示不存在；SC_ESC:\"\" 表示原始空字符串。", "单侧缺失金额/差额为空，真实零写 0 分与 0.00。")
    yield ("安全", "raw_numeric_values", "原始数值按读取值转文本并另列类型。", "日期为 ISO 文本；非有限数值可显示 inf/nan，仅作为问题行，不参与金额计算。")
    yield ("安全", "static_report", "无公式、超链接、外部链接或宏；汇总均为引擎静态值。", "不执行公式、不刷新外链。修正源文件后重新核对，不修改本报告作为新来源。")


def _source_fields(entry, table):
    if entry is None:
        return {key: None for key in ("source_id", "filename", "sheet", "row", "physical_start", "physical_end", "id_column", "amount_column")}
    pos = entry.row.source.position
    return {"source_id": entry.source_id, "filename": encode_text(table.source.spec.filename),
            "sheet": encode_text(pos.sheet) if pos.sheet is not None else None, "row": pos.row,
            "physical_start": pos.physical_start, "physical_end": pos.physical_end,
            "id_column": table.mapping.id_column, "amount_column": table.mapping.amount_column}


def _business_rows(result, entries, exceptions=False):
    for record in result.records:
        if exceptions and record.status == Status.EXACT_MATCH:
            continue
        data = {"status": str(record.status), "id": encode_text(record.id), "result_id": record.result_id,
                "order_minor": record.order_minor, "payment_minor": record.payment_minor, "delta_minor": record.delta_minor,
                "order_amount": _money(record.order_minor), "payment_amount": _money(record.payment_minor), "delta_amount": _money(record.delta_minor)}
        for side, source_id, table in (("order", record.order_source_id, result.tables[0]), ("payment", record.payment_source_id, result.tables[1])):
            data.update({f"{side}_{key}": value for key, value in _source_fields(entries.get(source_id), table).items()})
        yield tuple(data[column.key] for column in BUSINESS_COLUMNS)


def _ledger_rows(result, issues=False):
    tables = {t.source.spec.side: t for t in result.tables}
    for entry in result.ledger:
        if issues and entry.status != Status.DATA_ISSUE:
            continue
        row, side = entry.row, entry.row.source.side
        table = tables[side]
        data = _source_fields(entry, table)
        data.update({"status": str(entry.status), "side": str(side), "id_raw": _raw_text(row.id_raw),
                     "amount_raw": _raw_text(row.amount_raw), "amount_display": _money(row.amount_minor),
                     "amount_minor": row.amount_minor, "problem_codes": _json(entry.problem_codes),
                     "normalization_notes": _json(row.normalization_notes), "result_id": entry.result_id,
                     "normalized_id": encode_text(row.normalized_id) if row.normalized_id is not None else None,
                     "id_type": row.raw_types[0], "amount_type": row.raw_types[1], "sha256": table.source.sha256,
                     "hidden": "是" if row.source.position.hidden else "否"})
        if issues:
            data["repair_suggestions"] = "\n".join(PROBLEM_HELP[code] for code in entry.problem_codes)
        yield tuple(data[c.key] for c in (ISSUE_COLUMNS if issues else LEDGER_COLUMNS))


def _rows(result, name, entries):
    if name == "运行说明":
        return _run_rows(result)
    if name in ("业务结果", "异常明细"):
        return _business_rows(result, entries, name == "异常明细")
    return _ledger_rows(result, name == "数据问题")


def _headers(name):
    return RUN_HEADERS if name == "运行说明" else tuple(c.title for c in _COLUMN_SETS[name])


class _MemoryExcelWriter(ExcelWriter):
    def write_worksheet(self, ws):
        # Standard openpyxl writes worksheet XML to disk even with a BytesIO
        # final target. Its pinned internal writer accepts an explicit buffer.
        ws._drawing = SpreadsheetDrawing()
        with BytesIO() as out:
            writer = WorksheetWriter(ws, out=out)
            try:
                writer.write()
                ws._rels = writer._rels
                self._archive.writestr(ws.path[1:], writer.read())
                self.manifest.append(ws)
            finally:
                writer.close()


def _serialize_workbook(workbook):
    with BytesIO() as buffer:
        with ZipFile(buffer, "w", ZIP_DEFLATED, allowZip64=True) as archive:
            _MemoryExcelWriter(workbook, archive).save()
        return buffer.getvalue()


def _add_row(sheet, values, header=False):
    sheet.append(values)
    number = sheet._current_row
    cells = [sheet.cell(number, index) for index in range(1, len(values) + 1)]
    for cell, value in zip(cells, values):
        if value is None:
            continue
        if isinstance(value, str):
            _text_length(value)  # catch system messages too; no silent truncation
            cell.data_type = "s"  # overrides openpyxl's =formula/#N/A inference
            cell.number_format = "@"
            cell.alignment = _TEXT_ALIGN
        elif type(value) is int:
            cell.number_format = "0"
            cell.alignment = _NUMBER_ALIGN
        else:
            raise ExportError("EXPORT_VALUE_TYPE", "报告字段类型不受支持，已中止。")
        cell.font = _BODY_FONT
        if header:
            cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    # Preserve lines and make routine row content legible without shrinking it.
    height = 22
    for cell, value in zip(cells, values):
        if isinstance(value, str):
            width = sheet.column_dimensions[cell.column_letter].width or 20
            visual_units = max((sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in line)
                                for line in value.split("\n")), default=0)
            lines = max(len(value.split("\n")), (visual_units + int(width) - 1) // max(int(width), 1))
            # Long diagnostic fields remain intact and inspectable in formula
            # bar; Excel itself limits row height to 409 points.
            height = max(height, min(409, 16 * lines + 8))
    sheet.row_dimensions[number].height = max(36, height) if header else height


def report_filename(result):
    try:
        run_id = result.manifest.run_id
        if not re.fullmatch(r"[a-f0-9]{32}", run_id):
            raise ValueError
        timestamp = datetime.fromisoformat(result.manifest.finished_at)
        if timestamp.tzinfo is None:
            raise ValueError
        timestamp = timestamp.astimezone(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        return f"sheetcheck_{timestamp}_{run_id}.xlsx"
    except (ValueError, TypeError, AttributeError):
        raise ExportError("EXPORT_MANIFEST", "运行标识或时间无效，无法生成安全报告名称。") from None


def _write_xlsx(result: ReconciliationResult) -> bytes:
    """Keep every writer/worksheet reference inside this completed scope."""
    workbook = Workbook()
    try:
        workbook.remove(workbook.active)
        workbook.properties.creator = "SheetCheck"
        workbook.properties.lastModifiedBy = "SheetCheck"
        workbook.properties.title = "SheetCheck 核对报告"
        workbook.calculation = CalcProperties(calcMode="manual", fullCalcOnLoad=False, forceFullCalc=False)
        entries = {e.source_id: e for e in result.ledger}
        for name in SHEET_NAMES:
            sheet = workbook.create_sheet(name)
            sheet.sheet_view.showGridLines = False
            sheet.freeze_panes = "C2" if name != "运行说明" else "A2"
            widths = (16, 45, 80, 75) if name == "运行说明" else tuple(c.width for c in _COLUMN_SETS[name])
            for index, width in enumerate(widths, 1):
                sheet.column_dimensions[get_column_letter(index)].width = width
            _add_row(sheet, _headers(name), header=True)
            for values in _rows(result, name, entries):
                _add_row(sheet, values)
            sheet.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{sheet.max_row}"
        data = _serialize_workbook(workbook)
    except (ExportError, IntegrityError):
        raise
    except Exception:
        raise ExportError("EXPORT_WRITE_FAILED", "报告写出失败，未产生成功报告；请重试或检查输入。") from None
    finally:
        workbook.close()
    return data


def export_xlsx(result: ReconciliationResult) -> bytes:
    assert_invariants(result)  # absolutely no workbook creation on invalid input
    report_filename(result)
    data = _write_xlsx(result)
    # Worksheets and cells point back to their workbook. Merely deleting the
    # workbook variable leaves both cycles and the last local worksheet alive.
    # The helper has now returned, so collect its unreachable object graph
    # before starting the mandatory streaming readback.
    gc.collect()
    audit_xlsx(data, result)
    return data


def _require(ok, code="REPORT_READBACK_MISMATCH"):
    if not ok:
        raise IntegrityError(code, "报告回读审计失败，未产生成功报告。")


def _package_audit(data, counts):
    _require(isinstance(data, bytes) and len(data) <= MAX_REPORT_BYTES, "REPORT_SIZE")
    base_parts = {"docProps/app.xml", "docProps/core.xml", "xl/theme/theme1.xml", "xl/styles.xml",
                  "_rels/.rels", "xl/workbook.xml", "xl/_rels/workbook.xml.rels", "[Content_Types].xml"}
    sheet_parts = {f"xl/worksheets/sheet{i}.xml": (name, counts[name], len(_headers(name))) for i, name in enumerate(SHEET_NAMES, 1)}
    with ZipFile(BytesIO(data)) as archive:
        infos = archive.infolist()
        _require(len(infos) == len(base_parts | sheet_parts.keys()) and {i.filename for i in infos} == base_parts | sheet_parts.keys(), "REPORT_PACKAGE_PARTS")
        _require(sum(i.file_size for i in infos) <= MAX_REPORT_XML_BYTES, "REPORT_EXPANDED_SIZE")
        for info in infos:
            _require(not info.flag_bits & 1 and info.compress_type in (0, ZIP_DEFLATED), "REPORT_PACKAGE_ENCODING")
            stack, row_number, last_column, seen_rows = [], 0, 0, 0
            relation_ids = set()
            filter_names = set()
            with archive.open(info) as stream:
                for event, element in iterparse(stream, events=("start", "end"), forbid_dtd=True, forbid_entities=True, forbid_external=True):
                    local = element.tag.rsplit("}", 1)[-1]
                    if event == "start":
                        stack.append(element)
                        _require(len(stack) <= 64, "REPORT_XML_DEPTH")
                        _require(local not in {"f", "formula", "formula1", "formula2", "conditionalFormatting", "cfRule",
                                              "hyperlink", "hyperlinks", "externalReference", "oleObject", "oleObjects",
                                              "controls", "dataValidations", "extLst", "AlternateContent"}, "REPORT_ACTIVE_CONTENT")
                        if len(stack) == 2 and info.filename in sheet_parts:
                            _require(local in {"sheetPr", "dimension", "sheetViews", "sheetFormatPr", "cols", "sheetData", "autoFilter", "pageMargins"}, "REPORT_SHEET_STRUCTURE")
                        if len(stack) == 2 and info.filename == "xl/workbook.xml":
                            _require(local in {"workbookPr", "workbookProtection", "bookViews", "sheets", "definedNames", "calcPr"}, "REPORT_WORKBOOK_STRUCTURE")
                        if local == "Relationship":
                            ident, target, kind = element.get("Id"), element.get("Target", ""), element.get("Type", "").rsplit("/", 1)[-1]
                            _require(ident and ident not in relation_ids, "REPORT_RELATIONSHIP")
                            relation_ids.add(ident)
                            _require(element.get("TargetMode", "Internal") == "Internal" and not re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*:", target) and not target.startswith("//"), "REPORT_EXTERNAL_LINK")
                            path = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join("xl" if info.filename.startswith("xl/") else "", target))
                            allowed = ({("officeDocument", "xl/workbook.xml"), ("core-properties", "docProps/core.xml"), ("extended-properties", "docProps/app.xml")}
                                       if info.filename == "_rels/.rels" else {("worksheet", p) for p in sheet_parts} | {("styles", "xl/styles.xml"), ("theme", "xl/theme/theme1.xml")})
                            _require((kind, path) in allowed, "REPORT_RELATIONSHIP")
                        if info.filename == "[Content_Types].xml":
                            content_type = element.get("ContentType", "").lower()
                            _require(not any(x in content_type for x in ("macro", "vba", "external", "activex", "ole", "connection")), "REPORT_ACTIVE_CONTENT")
                        if info.filename in sheet_parts:
                            _, max_rows, max_columns = sheet_parts[info.filename]
                            if local == "row":
                                row_number = int(element.get("r", "0"))
                                seen_rows += 1
                                _require(row_number == seen_rows and row_number <= max_rows, "REPORT_ROW_COORDINATES")
                                _require(element.get("hidden", "0") in ("0", "false"), "REPORT_HIDDEN_CONTENT")
                                last_column = 0
                            elif local == "c":
                                coordinate = element.get("r", "")
                                _require(bool(re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]*", coordinate)), "REPORT_CELL_COORDINATES")
                                row, col = coordinate_to_tuple(coordinate)
                                _require(row == row_number and last_column < col <= max_columns, "REPORT_CELL_COORDINATES")
                                last_column = col
                            elif local in {"mergeCell", "col"}:
                                _require(local != "mergeCell" and element.get("hidden", "0") in ("0", "false"), "REPORT_HIDDEN_CONTENT")
                    else:
                        if local == "definedName":
                            # Excel autofilters generate a static built-in range
                            # name. Permit only our exact bounded range for each
                            # fixed sheet, never arbitrary names or expressions.
                            index = int(element.get("localSheetId", "-1"))
                            _require(info.filename == "xl/workbook.xml" and 0 <= index < len(SHEET_NAMES), "REPORT_DEFINED_NAME")
                            name = SHEET_NAMES[index]
                            expected_range = f"'{name}'!$A$1:${get_column_letter(len(_headers(name)))}${counts[name]}"
                            _require(element.attrib == {"name": "_xlnm._FilterDatabase", "localSheetId": str(index), "hidden": "1"}
                                     and element.text == expected_range and index not in filter_names, "REPORT_DEFINED_NAME")
                            filter_names.add(index)
                        stack.pop()
                        if stack:
                            stack[-1].remove(element)
                        element.clear()
            if info.filename in sheet_parts:
                _require(seen_rows == sheet_parts[info.filename][1], "REPORT_ROW_COUNT")
            if info.filename == "xl/workbook.xml":
                _require(filter_names == set(range(len(SHEET_NAMES))), "REPORT_DEFINED_NAME")


def audit_xlsx(data: bytes, result: ReconciliationResult) -> None:
    """Reject unsafe or semantically changed saved bytes before release to caller."""
    assert_invariants(result)
    workbook = None
    try:
        entries = {e.source_id: e for e in result.ledger}
        counts = {name: 1 + sum(1 for _ in _rows(result, name, entries)) for name in SHEET_NAMES}
        _package_audit(data, counts)
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=True)
        _require(tuple(workbook.sheetnames) == SHEET_NAMES, "REPORT_SHEETS")
        _require(not workbook.defined_names and not workbook._external_links and workbook.vba_archive is None, "REPORT_ACTIVE_CONTENT")
        sentinel = object()
        for name in SHEET_NAMES:
            sheet = workbook[name]
            _require(sheet.sheet_state == "visible", "REPORT_HIDDEN_CONTENT")
            _require(sheet.max_row == counts[name] and sheet.max_column == len(_headers(name)), "REPORT_DIMENSIONS")
            expected = iter((_headers(name),))
            from itertools import chain
            expected = chain(expected, _rows(result, name, entries))
            for actual, values in zip_longest(sheet.iter_rows(), expected, fillvalue=sentinel):
                _require(actual is not sentinel and values is not sentinel and len(actual) == len(values))
                for cell, value in zip(actual, values):
                    _require(cell.data_type not in {"f", "e", "b", "d"}, "REPORT_CELL_TYPE")
                    _require(cell.value == value, "REPORT_CELL_VALUE")
                    if value is None:
                        _require(cell.value is None)
                    elif isinstance(value, str):
                        _require(cell.data_type == "s" and cell.number_format == "@", "REPORT_TEXT_TYPE")
                    else:
                        _require(type(cell.value) is int and cell.data_type == "n" and cell.number_format == "0", "REPORT_INTEGER_TYPE")
    except (IntegrityError, ExportError):
        raise
    except Exception:
        raise IntegrityError("REPORT_INVALID", "报告损坏或回读失败，未产生成功报告。") from None
    finally:
        if workbook is not None:
            workbook.close()

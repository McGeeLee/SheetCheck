"""Tamper the serialized XLSX, independently of the exporter's layout builder.

These are adversarial boundary tests: unchanged engine results must not bless a
workbook whose package, cells, ownership, or source locations have changed.
All source data and workbook mutations are synthetic and remain in memory.
"""
from __future__ import annotations

import csv
from dataclasses import replace
from io import BytesIO, StringIO
import warnings
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook

import sheetcheck.report as report
from sheetcheck.engine import reconcile
from sheetcheck.errors import SheetCheckError
from sheetcheck.models import FileSpec, Mapping, Rules, Side
from sheetcheck.readers import read_source
from sheetcheck.validators import validate


MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/package/2006/relationships"
DOCREL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
Q = lambda local: f"{{{MAIN}}}{local}"


def csv_bytes(rows):
    out = StringIO(newline="")
    csv.writer(out).writerows(rows)
    return out.getvalue().encode("utf-8")


@pytest.fixture(scope="module")
def report_result():
    rules = Rules(currency="CNY", period="2026-10", same_currency_confirmed=True, same_period_confirmed=True)
    tables = []
    for side, rows in (
        (Side.ORDER, [("00123", "12.30"), ("差额", "20.00"), ("单侧零", "0.00"), ("坏行", "BAD"), ("", "")]),
        (Side.PAYMENT, [("00123", "12.30"), ("差额", "19.00"), ("仅收款", "1.00"), ("坏行", "2.00")]),
    ):
        source = read_source(csv_bytes([("编号", "金额"), *rows]), FileSpec(f"{side}.csv", side))
        tables.append(validate(source, Mapping(1, 2), rules))
    return reconcile(*tables, rules)


@pytest.fixture(scope="module")
def report_bytes(report_result):
    return report.export_xlsx(report_result)


def archive_edit(data, changes=None, *, remove=(), additions=None):
    """Change raw package members; no workbook library normalizes the attack."""
    changes, additions = changes or {}, additions or {}
    output = BytesIO()
    with ZipFile(BytesIO(data)) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for info in source.infolist():
            if info.filename in remove:
                continue
            value = source.read(info.filename)
            if info.filename in changes:
                value = changes[info.filename](value)
            target.writestr(info.filename, value)
        for name, value in additions.items():
            target.writestr(name, value)
    return output.getvalue()


def xml_edit(data, member, edit):
    def mutate(raw):
        root = ET.fromstring(raw)
        edit(root)
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)
    return archive_edit(data, {member: mutate})


def cell_text(cell):
    return "".join(item.text or "" for item in cell.iter(Q("t")))


def find_text_cell(root, value):
    return next(cell for cell in root.iter(Q("c")) if cell_text(cell) == value)


def set_text(cell, value):
    for child in list(cell):
        cell.remove(child)
    cell.set("t", "inlineStr")
    ET.SubElement(ET.SubElement(cell, Q("is")), Q("t")).text = value


def target_cell(root, row_identifier, column_title):
    """Locate by documented labels, never by the exporter's expected-row code."""
    column_cell = find_text_cell(root, column_title)
    column = "".join(ch for ch in column_cell.get("r") if ch.isalpha())
    identity_cell = find_text_cell(root, row_identifier)
    row = next(row for row in root.iter(Q("row")) if identity_cell in list(row))
    ref = column + row.get("r")
    existing = next((cell for cell in row if cell.get("r") == ref), None)
    return existing if existing is not None else ET.SubElement(row, Q("c"), {"r": ref})


def set_number(cell, value, kind="n"):
    for child in list(cell):
        cell.remove(child)
    cell.set("t", kind)
    ET.SubElement(cell, Q("v")).text = str(value)


def require_rejected(data, result):
    with pytest.raises(SheetCheckError):
        report.audit_xlsx(data, result)


@pytest.mark.parametrize("sheet_number", range(1, 6))
def test_readback_rejects_formula_in_every_sheet(report_bytes, report_result, sheet_number):
    def inject(root):
        cell = next(root.iter(Q("c")))
        for child in list(cell):
            cell.remove(child)
        cell.attrib.pop("t", None)
        ET.SubElement(cell, Q("f")).text = 'HYPERLINK("https://invalid.example/", "click")'
        ET.SubElement(cell, Q("v")).text = "0"
    data = xml_edit(report_bytes, f"xl/worksheets/sheet{sheet_number}.xml", inject)
    require_rejected(data, report_result)


def test_readback_rejects_internal_hyperlink_without_relationship(report_bytes, report_result):
    def inject(root):
        links = ET.SubElement(root, Q("hyperlinks"))
        ET.SubElement(links, Q("hyperlink"), {"ref": "A1", "location": "'业务结果'!A1"})
    require_rejected(xml_edit(report_bytes, "xl/worksheets/sheet1.xml", inject), report_result)


def test_readback_rejects_conditional_formatting_formula(report_bytes, report_result):
    def inject(root):
        formatting = ET.SubElement(root, Q("conditionalFormatting"), {"sqref": "A2"})
        rule = ET.SubElement(formatting, Q("cfRule"), {"type": "expression", "priority": "1"})
        ET.SubElement(rule, Q("formula")).text = "1+1=2"
    require_rejected(xml_edit(report_bytes, "xl/worksheets/sheet2.xml", inject), report_result)


@pytest.mark.parametrize("target_mode", ["External", "Internal"])
def test_readback_rejects_hyperlink_relationship_even_if_orphaned(report_bytes, report_result, target_mode):
    root = ET.Element(f"{{{REL}}}Relationships")
    ET.SubElement(root, f"{{{REL}}}Relationship", {
        "Id": "rIdInjected", "Type": f"{DOCREL}/hyperlink",
        "Target": "https://invalid.example/" if target_mode == "External" else "../worksheets/sheet2.xml",
        "TargetMode": target_mode,
    })
    data = archive_edit(report_bytes, additions={"xl/worksheets/_rels/sheet1.xml.rels": ET.tostring(root)})
    require_rejected(data, report_result)


@pytest.mark.parametrize("part", [
    "xl/vbaProject.bin",
    "xl/externalLinks/externalLink1.xml",
    "xl/embeddings/oleObject1.bin",
    "xl/activeX/activeX1.bin",
    "xl/connections.xml",
])
def test_readback_rejects_active_content_parts_even_when_unreferenced(report_bytes, report_result, part):
    data = archive_edit(report_bytes, additions={part: b"synthetic-active-part"})
    require_rejected(data, report_result)


def test_readback_rejects_external_non_hyperlink_relationship(report_bytes, report_result):
    def inject(root):
        ET.SubElement(root, f"{{{REL}}}Relationship", {
            "Id": "rIdInjected", "Type": f"{DOCREL}/image",
            "Target": "https://invalid.example/image.png", "TargetMode": "External",
        })
    require_rejected(xml_edit(report_bytes, "xl/_rels/workbook.xml.rels", inject), report_result)


@pytest.mark.parametrize("mutation", ["expression", "name", "range", "duplicate"])
def test_readback_allows_only_exact_static_filter_defined_names(report_bytes, report_result, mutation):
    def inject(root):
        name = next(root.iter(Q("definedName")))
        if mutation == "expression":
            name.text = "1+1"
        elif mutation == "name":
            name.set("name", "attacker_defined_name")
        elif mutation == "range":
            name.text = "'业务结果'!$A$1:$XFD$1048576"
        else:
            names = root.find(Q("definedNames"))
            names.append(ET.fromstring(ET.tostring(name)))
    require_rejected(xml_edit(report_bytes, "xl/workbook.xml", inject), report_result)


@pytest.mark.parametrize("mutation", ["extra", "missing", "reordered", "renamed", "hidden"])
def test_readback_rejects_changed_sheet_set_and_order(report_bytes, report_result, mutation):
    workbook = load_workbook(BytesIO(report_bytes))
    if mutation == "extra":
        workbook.create_sheet("附加页")
    elif mutation == "missing":
        workbook.remove(workbook.worksheets[-1])
    elif mutation == "reordered":
        workbook.move_sheet(workbook.worksheets[-1], offset=-4)
    elif mutation == "renamed":
        workbook.worksheets[0].title = "伪造运行说明"
    else:
        workbook.worksheets[-1].sheet_state = "hidden"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    require_rejected(output.getvalue(), report_result)


@pytest.mark.parametrize("mutation", ["extra_row", "extra_column", "missing_row", "duplicate_row"])
def test_readback_rejects_changed_data_extent(report_bytes, report_result, mutation):
    def mutate(root):
        data = root.find(Q("sheetData"))
        rows = list(data)
        if mutation == "extra_row":
            row = ET.SubElement(data, Q("row"), {"r": "999"})
            set_text(ET.SubElement(row, Q("c"), {"r": "A999"}), "injected")
        elif mutation == "extra_column":
            set_text(ET.SubElement(rows[-1], Q("c"), {"r": "XFD" + rows[-1].get("r")}), "injected")
        elif mutation == "missing_row":
            data.remove(rows[-1])
        else:
            data.append(ET.fromstring(ET.tostring(rows[-1])))
    require_rejected(xml_edit(report_bytes, "xl/worksheets/sheet2.xml", mutate), report_result)


@pytest.mark.parametrize("mutation", ["text_value", "numeric_type", "duplicate_cell", "bad_reference"])
def test_readback_rejects_identifier_tampering(report_bytes, report_result, mutation):
    def mutate(root):
        cell = find_text_cell(root, "00123")
        if mutation == "text_value":
            set_text(cell, "123")
        elif mutation == "numeric_type":
            for child in list(cell):
                cell.remove(child)
            cell.set("t", "n")
            ET.SubElement(cell, Q("v")).text = "123"
        elif mutation == "duplicate_cell":
            row = next(row for row in root.iter(Q("row")) if cell in list(row))
            row.append(ET.fromstring(ET.tostring(cell)))
        else:
            cell.set("r", "A0")
    require_rejected(xml_edit(report_bytes, "xl/worksheets/sheet2.xml", mutate), report_result)


def test_readback_rejects_changed_source_id(report_bytes, report_result):
    source_id = report_result.ledger[0].source_id
    def mutate(root):
        set_text(find_text_cell(root, source_id), "forged-source-id")
    require_rejected(xml_edit(report_bytes, "xl/worksheets/sheet5.xml", mutate), report_result)


@pytest.mark.parametrize("column", ["订单金额（分）", "收款金额（分）", "差额（分）"])
def test_readback_rejects_changed_integer_amount(report_bytes, report_result, column):
    def mutate(root):
        set_number(target_cell(root, "差额", column), 777)
    require_rejected(xml_edit(report_bytes, "xl/worksheets/sheet2.xml", mutate), report_result)


@pytest.mark.parametrize("column", ["收款金额（分）", "差额（分）"])
def test_readback_does_not_accept_zero_instead_of_missing(report_bytes, report_result, column):
    def mutate(root):
        set_number(target_cell(root, "单侧零", column), 0)
    require_rejected(xml_edit(report_bytes, "xl/worksheets/sheet2.xml", mutate), report_result)


def test_readback_does_not_accept_missing_instead_of_true_zero(report_bytes, report_result):
    def mutate(root):
        cell = target_cell(root, "单侧零", "订单金额（分）")
        for child in list(cell):
            cell.remove(child)
        cell.attrib.pop("t", None)
    require_rejected(xml_edit(report_bytes, "xl/worksheets/sheet2.xml", mutate), report_result)


@pytest.mark.parametrize("kind", ["text", "boolean"])
def test_readback_rejects_same_amount_with_wrong_cell_type(report_bytes, report_result, kind):
    def mutate(root):
        cell = target_cell(root, "单侧零", "订单金额（分）")
        set_text(cell, "0") if kind == "text" else set_number(cell, 0, "b")
    require_rejected(xml_edit(report_bytes, "xl/worksheets/sheet2.xml", mutate), report_result)


@pytest.mark.parametrize("column", ["订单源行 / CSV逻辑记录", "订单CSV物理起始行", "订单CSV物理结束行", "订单编号列号", "订单金额列号"])
def test_readback_rejects_source_position_and_mapping_tampering(report_bytes, report_result, column):
    def mutate(root):
        set_number(target_cell(root, "00123", column), 999)
    require_rejected(xml_edit(report_bytes, "xl/worksheets/sheet2.xml", mutate), report_result)


def test_readback_rejects_duplicated_zip_member(report_bytes, report_result):
    with ZipFile(BytesIO(report_bytes)) as source:
        duplicate = source.read("xl/worksheets/sheet2.xml")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        data = archive_edit(report_bytes, additions={"xl/worksheets/sheet2.xml": duplicate})
    require_rejected(data, report_result)


@pytest.mark.parametrize("data", [b"", b"not-an-xlsx", b"PK\x03\x04truncated"])
def test_readback_converts_corrupt_container_to_safe_error(report_result, data):
    require_rejected(data, report_result)


def test_readback_rejects_corrupt_sheet_xml(report_bytes, report_result):
    data = archive_edit(report_bytes, {"xl/worksheets/sheet2.xml": lambda raw: b"<worksheet>"})
    require_rejected(data, report_result)


def test_readback_rejects_xml_entity_declaration(report_bytes, report_result):
    def inject(raw):
        declaration = b'<!DOCTYPE worksheet [<!ENTITY synthetic "not-source-data">]>'
        if raw.startswith(b"<?xml"):
            head, rest = raw.split(b"?>", 1)
            return head + b"?>" + declaration + rest
        return declaration + raw
    data = archive_edit(report_bytes, {"xl/worksheets/sheet2.xml": inject})
    require_rejected(data, report_result)


def test_export_stops_when_kernel_invariant_fails(report_result):
    invalid = replace(report_result, ledger=report_result.ledger[1:])
    with pytest.raises(SheetCheckError):
        report.export_xlsx(invalid)


def test_export_calls_readback_gate_before_returning_bytes(report_result, monkeypatch):
    calls = []
    def reject(data, result):
        calls.append((data, result))
        raise SheetCheckError("TEST_READBACK_REJECTED", "synthetic rejection")
    monkeypatch.setattr(report, "audit_xlsx", reject)
    with pytest.raises(SheetCheckError):
        report.export_xlsx(report_result)
    assert len(calls) == 1
    assert isinstance(calls[0][0], bytes)
    assert calls[0][1] is report_result


def test_export_does_not_return_workbook_corrupted_during_save(report_result, monkeypatch):
    original = report._serialize_workbook
    def corrupt_after_save(workbook):
        output = original(workbook)
        def inject(root):
            links = ET.SubElement(root, Q("hyperlinks"))
            ET.SubElement(links, Q("hyperlink"), {"ref": "A1", "location": "A1"})
        return xml_edit(output, "xl/worksheets/sheet1.xml", inject)
    monkeypatch.setattr(report, "_serialize_workbook", corrupt_after_save)
    with pytest.raises(SheetCheckError):
        report.export_xlsx(report_result)

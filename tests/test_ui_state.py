"""UI-independent session transitions and metadata-import security contracts."""
from dataclasses import replace
from contextlib import contextmanager
from io import BytesIO
from threading import Event, Thread, current_thread
import zipfile

from openpyxl import Workbook
import pytest

from sheetcheck.errors import ExportError, InputError, IntegrityError
from sheetcheck.models import FileSpec, Limits, Mapping, Rules, Side
from sheetcheck.readers import list_worksheets
from sheetcheck.report import audit_xlsx
from sheetcheck.ui_state import InputSelection, Stage, Workflow
import sheetcheck.ui_state as state


DATA = b"id,amount,alternate\nA,1.00,1.00\n"
RULES = Rules(period="fictional", same_currency_confirmed=True, same_period_confirmed=True)


def selections(data=DATA):
    return {side: InputSelection(data, FileSpec(f"{side}.csv", side)) for side in Side}


def imported():
    workflow = Workflow()
    workflow.sync_inputs(selections())
    for side in Side:
        workflow.get_source(side)
    return workflow


def configured():
    workflow = imported()
    workflow.configure({side: Mapping(1, 2) for side in Side}, RULES)
    return workflow


def finished():
    workflow = configured()
    workflow.run()
    workflow.prepare_report()
    return workflow


def assert_no_success(workflow):
    assert workflow.result is None
    assert workflow.report_bytes is None
    assert workflow.report_name is None
    assert workflow.stage not in {Stage.RECONCILED, Stage.EXPORTED}


def test_complete_workflow_commits_only_after_each_stage(monkeypatch):
    workflow = Workflow()
    assert workflow.stage == Stage.EMPTY
    assert not workflow.ready
    assert workflow.sync_inputs(selections())
    assert workflow.stage == Stage.EMPTY
    workflow.get_source(Side.ORDER)
    assert workflow.stage == Stage.IMPORTED
    workflow.get_source(Side.PAYMENT)
    assert workflow.configure({side: Mapping(1, 2) for side in Side}, RULES)
    assert workflow.stage == Stage.MAPPED
    assert workflow.ready
    original_reconcile = state.reconcile

    def observe_validation(*args):
        assert workflow.stage == Stage.VALIDATED
        assert_no_success(workflow)
        return original_reconcile(*args)

    monkeypatch.setattr(state, "reconcile", observe_validation)
    result = workflow.run()
    assert workflow.result is result
    assert workflow.stage == Stage.RECONCILED
    assert result.manifest.status_counts["EXACT_MATCH"] == 1
    report = workflow.prepare_report()
    assert report.startswith(b"PK")
    assert workflow.report_bytes is report
    assert workflow.report_name.endswith(f"_{result.manifest.run_id}.xlsx")
    assert workflow.stage == Stage.EXPORTED


def test_unchanged_inputs_and_configuration_keep_current_result_and_parse_once(monkeypatch):
    workflow = finished()
    result, report = workflow.result, workflow.report_bytes

    def forbidden(*args):
        raise AssertionError("unchanged data must reuse this session's cache")

    monkeypatch.setattr(state, "read_source", forbidden)
    monkeypatch.setattr(state, "export_xlsx", forbidden)
    assert not workflow.sync_inputs(selections())
    assert not workflow.configure({side: Mapping(1, 2) for side in Side}, RULES)
    for side in Side:
        assert workflow.get_source(side) is workflow.sources[side]
    assert workflow.result is result
    assert workflow.prepare_report() is report
    assert workflow.stage == Stage.EXPORTED


@pytest.mark.parametrize("change", [
    {"filename": "renamed.csv"}, {"sheet": "a different sheet"}, {"header_row": 2},
    {"encoding": "gb18030"}, {"delimiter": ";"},
    {"limits": replace(Limits(), max_records=10)},
    {"limits": replace(Limits(), max_file_bytes=1000)},
    {"limits": replace(Limits(), max_columns=3)},
    {"limits": replace(Limits(), max_zip_entries=10)},
    {"limits": replace(Limits(), max_uncompressed_bytes=1000)},
    {"limits": replace(Limits(), max_compression_ratio=100)},
])
def test_every_read_parameter_change_revokes_result_and_only_changed_cache(change):
    workflow = finished()
    untouched = workflow.sources[Side.PAYMENT]
    updated = dict(workflow.selections)
    current = updated[Side.ORDER]
    updated[Side.ORDER] = replace(current, spec=replace(current.spec, **change))
    assert workflow.sync_inputs(updated)
    assert_no_success(workflow)
    assert Side.ORDER not in workflow.sources
    assert workflow.mappings[Side.ORDER] is None
    assert workflow.sources[Side.PAYMENT] is untouched
    assert workflow.mappings[Side.PAYMENT] == Mapping(1, 2)


def test_same_filename_different_bytes_and_removal_are_changes():
    workflow = finished()
    updated = dict(workflow.selections)
    updated[Side.ORDER] = replace(updated[Side.ORDER], data=DATA.replace(b"1.00", b"2.00"))
    assert workflow.sync_inputs(updated)
    assert_no_success(workflow)
    source = workflow.get_source(Side.ORDER)
    assert source.rows[0].cells[1].value == "2.00"
    workflow.configure({side: Mapping(1, 2) for side in Side}, RULES)
    assert workflow.run().records[0].delta_minor == -100
    workflow.prepare_report()
    assert workflow.sync_inputs({Side.PAYMENT: updated[Side.PAYMENT]})
    assert workflow.selections[Side.ORDER] is None
    assert_no_success(workflow)
    assert not workflow.ready


@pytest.mark.parametrize("change", [
    {"currency": "USD"}, {"period": "next period"}, {"period": ""},
    {"same_currency_confirmed": False}, {"same_period_confirmed": False},
    {"version": "v2"}, {"tolerance_minor": 1}, {"minor_digits": 3},
    {"max_amount_minor": 10},
])
def test_every_rule_change_invalidates_result_and_keeps_parsed_sources(change):
    workflow = finished()
    sources = dict(workflow.sources)
    assert workflow.configure(workflow.mappings, replace(RULES, **change))
    assert_no_success(workflow)
    assert all(workflow.sources[side] is sources[side] for side in Side)


@pytest.mark.parametrize("mapping", [Mapping(1, 3), Mapping(2, 1), None, Mapping(1, 1), Mapping(0, 2)])
def test_mapping_change_invalidates_all_derived_outputs(mapping):
    workflow = finished()
    updated = dict(workflow.mappings)
    updated[Side.ORDER] = mapping
    assert workflow.configure(updated, RULES)
    assert_no_success(workflow)
    assert workflow.sources


@pytest.mark.parametrize("mapping", [None, Mapping(1, 1), Mapping(1, 4), Mapping(0, 2), Mapping(True, 2)])
def test_ready_requires_complete_valid_distinct_columns(mapping):
    workflow = imported()
    workflow.configure({Side.ORDER: mapping, Side.PAYMENT: Mapping(1, 2)}, RULES)
    assert not workflow.ready
    with pytest.raises(InputError, match="MAPPING_REQUIRED"):
        workflow.run()
    assert_no_success(workflow)


@pytest.mark.parametrize("rules", [Rules(), replace(RULES, period=" \n"),
    replace(RULES, same_currency_confirmed=False), replace(RULES, same_period_confirmed=False),
    replace(RULES, currency="JPY"), replace(RULES, tolerance_minor=1)])
def test_ready_requires_explicit_supported_rules(rules):
    workflow = configured()
    workflow.configure(workflow.mappings, rules)
    assert not workflow.ready
    with pytest.raises(InputError):
        workflow.run()
    assert_no_success(workflow)


def test_parse_failure_preserves_bytes_clears_stale_source_and_is_retryable(monkeypatch):
    workflow = finished()
    updated = dict(workflow.selections)
    updated[Side.ORDER] = replace(updated[Side.ORDER], data=b"id,amount,alternate\nB,2.00,2.00\n")
    workflow.sync_inputs(updated)
    read = state.read_source

    def fail_once(*args):
        raise InputError("CSV_PARSE_ERROR", "请修复 CSV。")

    monkeypatch.setattr(state, "read_source", fail_once)
    with pytest.raises(InputError, match="CSV_PARSE_ERROR"):
        workflow.get_source(Side.ORDER)
    assert workflow.selections[Side.ORDER] == updated[Side.ORDER]
    assert Side.ORDER not in workflow.sources
    assert Side.PAYMENT in workflow.sources
    assert_no_success(workflow)
    monkeypatch.setattr(state, "read_source", read)
    assert workflow.get_source(Side.ORDER).rows[0].cells[0].value == "B"


def test_missing_input_and_invalid_role_do_not_return_cached_success():
    workflow = finished()
    updated = dict(workflow.selections)
    updated[Side.ORDER] = updated[Side.PAYMENT]
    with pytest.raises(InputError, match="INPUT_SELECTION_INVALID"):
        workflow.sync_inputs(updated)
    assert_no_success(workflow)
    assert workflow.selections[Side.ORDER] is None
    with pytest.raises(InputError, match="INPUT_REQUIRED"):
        workflow.get_source(Side.ORDER)


@pytest.mark.parametrize("selection", ["not a selection", InputSelection(bytearray(DATA), FileSpec("file.csv", Side.ORDER)), InputSelection(DATA, None)])
def test_malformed_selection_revokes_old_source_and_report(selection):
    workflow = finished()
    updated = dict(workflow.selections)
    updated[Side.ORDER] = selection
    with pytest.raises(InputError, match="INPUT_SELECTION_INVALID"):
        workflow.sync_inputs(updated)
    assert_no_success(workflow)
    assert Side.ORDER not in workflow.sources
    assert not workflow.ready


@pytest.mark.parametrize("mappings,rules", [
    ({Side.ORDER: "bad", Side.PAYMENT: Mapping(1, 2)}, RULES),
    ({side: Mapping(1, 2) for side in Side}, "bad"),
])
def test_malformed_configuration_cannot_reuse_old_valid_configuration(mappings, rules):
    workflow = finished()
    with pytest.raises(InputError, match="CONFIGURATION_INVALID"):
        workflow.configure(mappings, rules)
    assert_no_success(workflow)
    assert not workflow.ready
    with pytest.raises(InputError):
        workflow.run()


@pytest.mark.parametrize("operation", ["validate", "reconcile"])
def test_failed_rerun_cannot_retain_previous_success_and_keeps_inputs(monkeypatch, operation):
    workflow = finished()
    original = getattr(state, operation)
    before = dict(workflow.selections)

    def fail(*args):
        raise IntegrityError("TEST_FAILURE", "测试失败。")

    monkeypatch.setattr(state, operation, fail)
    with pytest.raises(IntegrityError, match="TEST_FAILURE"):
        workflow.run()
    assert_no_success(workflow)
    assert workflow.selections == before
    assert len(workflow.sources) == 2
    monkeypatch.setattr(state, operation, original)
    assert workflow.run().records[0].id == "A"


@pytest.mark.parametrize("operation", ["export_xlsx", "report_filename"])
def test_export_failure_exposes_no_partial_report_and_preserves_retry_state(monkeypatch, operation):
    workflow = configured()
    result = workflow.run()
    original = getattr(state, operation)

    def fail(*args):
        raise ExportError("TEST_EXPORT_FAILURE", "请重试导出。")

    monkeypatch.setattr(state, operation, fail)
    with pytest.raises(ExportError, match="TEST_EXPORT_FAILURE"):
        workflow.prepare_report()
    assert workflow.result is result
    assert workflow.report_bytes is None
    assert workflow.report_name is None
    assert workflow.stage == Stage.RECONCILED
    monkeypatch.setattr(state, operation, original)
    assert workflow.prepare_report().startswith(b"PK")


def test_prepare_report_without_current_result_is_illegal():
    workflow = finished()
    workflow.invalidate()
    with pytest.raises(InputError, match="RESULT_REQUIRED"):
        workflow.prepare_report()
    assert_no_success(workflow)
    assert workflow.ready
    assert workflow.stage == Stage.MAPPED


def test_invalidation_runs_immediately_before_input_values_are_synchronized():
    workflow = finished()
    before = dict(workflow.selections)
    workflow.invalidate()
    assert_no_success(workflow)
    assert workflow.selections == before
    assert len(workflow.sources) == 2
    assert not workflow.sync_inputs(before)
    assert not workflow.configure(workflow.mappings, RULES)
    assert_no_success(workflow)


def test_reset_clears_original_instance_and_new_session_does_not_share_state():
    workflow = finished()
    other = Workflow()
    alias = workflow
    workflow.reset()
    assert alias is workflow
    assert alias.stage == Stage.EMPTY
    assert alias.selections == {side: None for side in Side}
    assert alias.mappings == {side: None for side in Side}
    assert alias.sources == {}
    assert alias.rules == Rules()
    assert_no_success(alias)
    assert not alias.ready
    workflow.sync_inputs(selections())
    assert other.selections == {side: None for side in Side}
    assert other.sources == {}
    assert workflow.selections is not other.selections
    assert workflow.mappings is not other.mappings


def workbook_bytes():
    workbook = Workbook()
    workbook.active.title = "订单"
    workbook.active.append(["id", "amount"])
    workbook.active.append(["A", "=1+1"])
    workbook.create_sheet("隐藏表").sheet_state = "hidden"
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def rewrite_zip(data, member, transform):
    output = BytesIO()
    with zipfile.ZipFile(BytesIO(data)) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as destination:
        for info in source.infolist():
            value = source.read(info.filename)
            destination.writestr(info.filename, transform(value) if info.filename == member else value)
    return output.getvalue()


def test_sheet_discovery_lists_hidden_names_without_reading_cells_or_executing_formulas(monkeypatch):
    import sheetcheck.readers as readers

    def no_cell_load(*args, **kwargs):
        raise AssertionError("metadata discovery must not load a workbook")

    monkeypatch.setattr(readers, "load_workbook", no_cell_load)
    assert list_worksheets(workbook_bytes(), FileSpec("book.xlsx", Side.ORDER)) == ("订单", "隐藏表")


@pytest.mark.parametrize("data,spec,code", [
    (DATA, FileSpec("book.csv", Side.ORDER), "XLSX_REQUIRED"),
    (b"not a workbook", FileSpec("book.xlsx", Side.ORDER), "FILE_SIGNATURE_INVALID"),
    (b"PK\x03\x04broken", FileSpec("book.xlsx", Side.ORDER), "XLSX_INVALID"),
    (b"", FileSpec("book.xlsx", Side.ORDER), "FILE_EMPTY"),
])
def test_sheet_discovery_rejects_invalid_inputs_with_safe_errors(data, spec, code):
    with pytest.raises(InputError, match=code):
        list_worksheets(data, spec)


@pytest.mark.parametrize("limits,code", [
    (replace(Limits(), max_file_bytes=1), "FILE_SIZE_LIMIT"),
    (replace(Limits(), max_zip_entries=1), "XLSX_ZIP_ENTRY_LIMIT"),
    (replace(Limits(), max_uncompressed_bytes=1), "XLSX_UNCOMPRESSED_LIMIT"),
    (replace(Limits(), max_compression_ratio=1), "XLSX_COMPRESSION_RATIO"),
])
def test_sheet_discovery_respects_all_container_limits(limits, code):
    with pytest.raises(InputError, match=code):
        list_worksheets(workbook_bytes(), FileSpec("book.xlsx", Side.ORDER, limits=limits))


def test_sheet_discovery_rejects_external_relationship_before_showing_names():
    data = rewrite_zip(workbook_bytes(), "xl/_rels/workbook.xml.rels", lambda value: value.replace(
        b"</Relationships>", b'<Relationship Id="evil" Type="http://example.invalid/hyperlink" Target="https://example.invalid/secret" TargetMode="External"/></Relationships>'
    ))
    with pytest.raises(InputError, match="XLSX_EXTERNAL_LINK"):
        list_worksheets(data, FileSpec("book.xlsx", Side.ORDER))


def test_sheet_discovery_preflights_unselected_sheet_xml():
    data = rewrite_zip(workbook_bytes(), "xl/worksheets/sheet2.xml", lambda value: b'<!DOCTYPE x [<!ENTITY e "bad">]><x>&e;</x>')
    with pytest.raises(InputError, match="XLSX_INVALID"):
        list_worksheets(data, FileSpec("book.xlsx", Side.ORDER))


def test_sheet_discovery_rejects_duplicate_sheet_names():
    data = rewrite_zip(workbook_bytes(), "xl/workbook.xml", lambda value: value.replace("隐藏表".encode(), "订单".encode()))
    with pytest.raises(InputError, match="XLSX_CONTAINER_INVALID"):
        list_worksheets(data, FileSpec("book.xlsx", Side.ORDER))


def test_sheet_discovery_rejects_missing_worksheet_relationship_target():
    data = rewrite_zip(workbook_bytes(), "xl/_rels/workbook.xml.rels", lambda value: value.replace(b"/xl/worksheets/sheet1.xml", b"/xl/worksheets/missing.xml"))
    with pytest.raises(InputError, match="XLSX_CONTAINER_INVALID"):
        list_worksheets(data, FileSpec("book.xlsx", Side.ORDER))


@contextmanager
def paused_operation(monkeypatch, function, action, *, fail=False):
    """Pause exactly the background operation, with deterministic handshakes."""
    entered, resume = Event(), Event()
    outcome = {"values": [], "errors": []}
    original = getattr(state, function)

    def controlled(*args, **kwargs):
        value = original(*args, **kwargs)
        if current_thread() is worker:
            entered.set()
            if not resume.wait(10):
                raise AssertionError("test did not resume the background operation")
            if fail:
                raise InputError("TEST_DELAYED_FAILURE", "旧操作失败。")
        return value

    def target():
        try:
            outcome["values"].append(action())
        except Exception as exc:
            outcome["errors"].append(exc)

    worker = Thread(target=target, daemon=True)
    with monkeypatch.context() as patch:
        patch.setattr(state, function, controlled)
        worker.start()
        try:
            assert entered.wait(10), "background operation did not reach its checkpoint"
            yield outcome
        finally:
            resume.set()
            worker.join(10)
            assert not worker.is_alive(), "background operation did not finish"


def assert_stale_operation(outcome):
    assert outcome["values"] == []
    assert len(outcome["errors"]) == 1
    assert isinstance(outcome["errors"][0], InputError)
    assert outcome["errors"][0].code == "OPERATION_STALE"


def replace_order(workflow):
    updated = dict(workflow.selections)
    updated[Side.ORDER] = replace(updated[Side.ORDER], data=DATA.replace(b"A,1.00", b"NEW,2.00"))
    workflow.sync_inputs(updated)
    workflow.get_source(Side.ORDER)


@pytest.mark.parametrize("fail", [False, True])
def test_old_parser_cannot_overwrite_new_source_or_erase_new_result(monkeypatch, fail):
    workflow = Workflow()
    workflow.sync_inputs(selections())
    with paused_operation(monkeypatch, "read_source", lambda: workflow.get_source(Side.ORDER), fail=fail) as outcome:
        replace_order(workflow)
        new_source = workflow.sources[Side.ORDER]
        workflow.get_source(Side.PAYMENT)
        workflow.configure({side: Mapping(1, 2) for side in Side}, RULES)
        result = workflow.run()
        report = workflow.prepare_report()
    assert_stale_operation(outcome)
    assert workflow.sources[Side.ORDER] is new_source
    assert workflow.result is result
    assert workflow.report_bytes is report
    assert workflow.stage == Stage.EXPORTED


@pytest.mark.parametrize("function", ["validate", "reconcile"])
@pytest.mark.parametrize("fail", [False, True])
def test_old_run_completion_or_failure_cannot_replace_new_input_result(monkeypatch, function, fail):
    workflow = configured()
    with paused_operation(monkeypatch, function, workflow.run, fail=fail) as outcome:
        replace_order(workflow)
        workflow.configure({side: Mapping(1, 2) for side in Side}, RULES)
        result = workflow.run()
        report = workflow.prepare_report()
    assert_stale_operation(outcome)
    assert workflow.result is result
    assert workflow.report_bytes is report
    assert workflow.stage == Stage.EXPORTED
    assert {record.id for record in result.records} == {"A", "NEW"}


@pytest.mark.parametrize("fail", [False, True])
def test_old_export_cannot_commit_bytes_or_filename_for_new_result(monkeypatch, fail):
    workflow = configured()
    old_result = workflow.run()
    with paused_operation(monkeypatch, "export_xlsx", workflow.prepare_report, fail=fail) as outcome:
        workflow.configure(workflow.mappings, replace(RULES, period="new-period"))
        result = workflow.run()
        report = workflow.prepare_report()
        filename = workflow.report_name
    assert_stale_operation(outcome)
    assert workflow.result is result
    assert workflow.report_bytes is report
    assert workflow.report_name == filename
    assert old_result.manifest.run_id not in filename
    assert result.manifest.run_id in filename
    assert workflow.stage == Stage.EXPORTED
    audit_xlsx(report, result)


@pytest.mark.parametrize("function", ["read_source", "validate", "reconcile", "export_xlsx"])
def test_reset_revokes_in_flight_work_and_old_object_stays_empty(monkeypatch, function):
    workflow = configured()
    if function == "read_source":
        workflow.reset()
        workflow.sync_inputs(selections())
        action = lambda: workflow.get_source(Side.ORDER)
    elif function == "export_xlsx":
        workflow.run()
        action = workflow.prepare_report
    else:
        action = workflow.run
    alias = workflow
    with paused_operation(monkeypatch, function, action) as outcome:
        workflow.reset()
    assert_stale_operation(outcome)
    assert_no_success(alias)
    assert alias.stage == Stage.EMPTY
    assert alias.sources == {}
    assert alias.selections == {side: None for side in Side}
    assert alias.rules == Rules()


@pytest.mark.parametrize("fail", [False, True])
def test_concurrent_read_of_same_input_only_newest_read_can_commit(monkeypatch, fail):
    workflow = Workflow()
    workflow.sync_inputs(selections())
    with paused_operation(monkeypatch, "read_source", lambda: workflow.get_source(Side.ORDER), fail=fail) as outcome:
        source = workflow.get_source(Side.ORDER)
    assert_stale_operation(outcome)
    assert workflow.sources[Side.ORDER] is source
    assert workflow.stage == Stage.IMPORTED


def test_repeated_run_on_same_inputs_supersedes_in_flight_run(monkeypatch):
    workflow = configured()
    with paused_operation(monkeypatch, "reconcile", workflow.run) as outcome:
        result = workflow.run()
        report = workflow.prepare_report()
    assert_stale_operation(outcome)
    assert workflow.result is result
    assert workflow.report_for_run(result.manifest.run_id) is report


def test_repeated_export_supersedes_in_flight_export_for_same_result(monkeypatch):
    workflow = configured()
    result = workflow.run()
    with paused_operation(monkeypatch, "export_xlsx", workflow.prepare_report) as outcome:
        report = workflow.prepare_report()
        filename = workflow.report_name
    assert_stale_operation(outcome)
    assert workflow.result is result
    assert workflow.report_for_run(result.manifest.run_id) is report
    assert workflow.report_name == filename


@pytest.mark.parametrize("change", ["invalidate", "reset", "rerun", "input", "rules"])
def test_atomic_download_getter_rejects_old_run_after_every_revocation(change):
    workflow = finished()
    old_id = workflow.result.manifest.run_id
    assert workflow.report_for_run(old_id) is workflow.report_bytes
    if change == "invalidate":
        workflow.invalidate()
    elif change == "reset":
        workflow.reset()
    elif change == "rerun":
        workflow.run()
        workflow.prepare_report()
    elif change == "input":
        replace_order(workflow)
    else:
        workflow.configure(workflow.mappings, replace(RULES, period="changed"))
    with pytest.raises(InputError, match="REPORT_STALE"):
        workflow.report_for_run(old_id)


def test_atomic_download_getter_rejects_missing_report_or_unrelated_run():
    workflow = configured()
    result = workflow.run()
    with pytest.raises(InputError, match="REPORT_STALE"):
        workflow.report_for_run(result.manifest.run_id)
    workflow.prepare_report()
    with pytest.raises(InputError, match="REPORT_STALE"):
        workflow.report_for_run("not-the-current-run")

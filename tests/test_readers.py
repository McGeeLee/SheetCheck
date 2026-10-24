"""Reader-level boundaries; all workbook and CSV data are synthetic."""
from __future__ import annotations

import io
import zipfile
from dataclasses import replace
from datetime import date

import pytest
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from sheetcheck.errors import InputError
from sheetcheck.models import FileSpec, Limits, Side
from sheetcheck.readers import read_source


def spec(extension="csv", **kwargs):
    return FileSpec(f"synthetic.{extension}", Side.ORDER, **kwargs)


def xlsx(make=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "订单"
    sheet.append(["id", "amount"])
    sheet.append(["00123", 12.3])
    if make:
        make(workbook, sheet)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def rewrite_archive(data, changes=None, additions=None):
    changes, additions = changes or {}, additions or {}
    result = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as original, zipfile.ZipFile(result, "w", zipfile.ZIP_DEFLATED) as target:
        for item in original.infolist():
            value = original.read(item.filename)
            if item.filename in changes:
                change = changes[item.filename]
                value = change(value) if callable(change) else change
            target.writestr(item.filename, value)
        for name, value in additions.items():
            target.writestr(name, value)
    return result.getvalue()


def assert_error(code, data, selected=None):
    with pytest.raises(InputError) as error:
        read_source(data, selected or spec())
    assert error.value.code == code


def test_csv_keeps_bom_leading_zeros_na_null_and_empty_fields():
    table = read_source(b"\xef\xbb\xbfid,amount\r\n00123,12.30\r\nNA,0\r\nNULL,0\r\n,5\r\n", spec())
    assert table.headers[0].value == "id"
    assert [row.cells[0].value for row in table.rows] == ["00123", "NA", "NULL", ""]
    assert all(cell.kind == "text" for row in table.rows for cell in row.cells)
    assert [(row.position.row, row.position.physical_start, row.position.physical_end) for row in table.rows] == [(2, 2, 2), (3, 3, 3), (4, 4, 4), (5, 5, 5)]


def test_csv_logical_records_differ_from_physical_lines():
    data = '说明\r\nid,amount,note\r\n"A,1",1,"first\r\nsecond"\r\n\r\nB,2,ok\r\n'.encode()
    table = read_source(data, spec(header_row=2))
    assert table.preamble_count == 1
    assert table.rows[0].cells[0].value == "A,1"
    assert table.rows[0].cells[2].value == "first\r\nsecond"
    assert [(row.position.row, row.position.physical_start, row.position.physical_end) for row in table.rows] == [(3, 3, 4), (4, 5, 5), (5, 6, 6)]
    assert [cell.value for cell in table.rows[1].cells] == ["", "", ""]


@pytest.mark.parametrize("delimiter", [",", "\t", ";"])
def test_csv_delimiters_and_gb18030_are_explicit(delimiter):
    data = f"编号{delimiter}金额\n中文{delimiter}1.00\n".encode("gb18030")
    assert_error("CSV_DECODE_ERROR", data)
    table = read_source(data, spec(encoding="gb18030", delimiter=delimiter))
    assert table.rows[0].cells[0].value == "中文"


@pytest.mark.parametrize("data", [b'id,amount\nA,"1\n', b'id,amount\nA,1,2\n', b'id,amount\nA\n', b'id,amount\nA"B,1\n', b'id,amount\n"A"x,1\n'])
def test_csv_malformed_input_blocks_import(data):
    with pytest.raises(InputError):
        read_source(data, spec())


def test_csv_duplicate_blank_headers_and_only_header():
    table = read_source(b"id,id,\n", spec())
    assert table.column_labels == ("1: id", "2: id", "3: ")
    assert table.rows == ()
    assert_error("FILE_EMPTY", b"")
    assert_error("HEADER_ROW_MISSING", b"\xef\xbb\xbf")
    assert_error("HEADER_ROW_MISSING", b"id,amount\n", spec(header_row=2))


def test_csv_all_empty_lines_are_not_discarded_or_invented():
    table = read_source(b"id,amount\n\n,\n", spec())
    assert len(table.rows) == 2
    assert all([cell.value for cell in row.cells] == ["", ""] for row in table.rows)
    assert len(read_source(b"id,amount\n", spec()).rows) == 0


def test_csv_source_ids_use_bytes_role_and_logical_position():
    data = b"id,amount\nA,1\nA,1\n"
    table = read_source(data, spec())
    repeated = read_source(data, spec())
    renamed = read_source(data, FileSpec("renamed.csv", Side.ORDER))
    payment = read_source(data, FileSpec("payment.csv", Side.PAYMENT))
    assert [row.source_id for row in table.rows] == [row.source_id for row in repeated.rows] == [row.source_id for row in renamed.rows]
    assert len({row.source_id for row in table.rows + payment.rows}) == 4
    changed = read_source(data + b"B,2\n", spec())
    assert table.rows[0].source_id != changed.rows[0].source_id


def test_input_extension_signatures_and_bytes_only():
    assert_error("FILE_TYPE_UNSUPPORTED", b"abc", spec("xls"))
    assert_error("FILE_TYPE_UNSUPPORTED", b"abc", spec("xlsm"))
    assert_error("INPUT_BYTES_REQUIRED", "id,amount\n")
    assert_error("FILE_SIGNATURE_INVALID", b"PK\x03\x04something")
    assert_error("FILE_SIGNATURE_INVALID", b"id,amount\n", spec("xlsx"))
    assert_error("FILE_SIGNATURE_INVALID", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", spec("xlsx"))
    assert_error("CSV_ENCODING_UNSUPPORTED", b"id,amount\n", spec(encoding="latin1"))
    assert_error("CSV_DELIMITER_UNSUPPORTED", b"id,amount\n", spec(delimiter="|"))
    assert_error("CSV_SHEET_INVALID", b"id,amount\n", spec(sheet="Sheet"))


def test_csv_capacity_boundaries_fail_without_truncation():
    data = b"id,amount\nA,1\nB,2\n"
    assert len(read_source(data, spec(limits=Limits(max_records=2))).rows) == 2
    assert_error("RECORD_LIMIT", data, spec(limits=Limits(max_records=1)))
    assert_error("COLUMN_LIMIT", data, spec(limits=Limits(max_columns=1)))
    assert_error("FILE_SIZE_LIMIT", data, spec(limits=Limits(max_file_bytes=len(data) - 1)))
    assert_error("PREAMBLE_LIMIT", data, spec(header_row=3, limits=Limits(max_records=1)))


def test_xlsx_preserves_types_hidden_rows_and_blank_gaps():
    def make(wb, ws):
        ws["A3"] = 123456789012345
        ws["B3"] = date(2026, 10, 23)
        ws["A4"] = True
        ws["B4"] = "#VALUE!"
        ws["B4"].data_type = "e"
        ws["A6"] = "NA"
        ws["B6"] = 0
        ws.row_dimensions[3].hidden = True
        ws["A1048576"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    table = read_source(xlsx(make), spec("xlsx"))
    assert table.spec.sheet == "订单"
    assert table.hidden_row_count == 1
    assert len(table.rows) == 5
    assert [cell.kind for cell in table.rows[0].cells] == ["text", "number"]
    assert [cell.kind for cell in table.rows[1].cells] == ["number", "date"]
    assert [cell.kind for cell in table.rows[2].cells] == ["boolean", "error"]
    assert [cell.kind for cell in table.rows[3].cells] == ["empty", "empty"]
    assert table.rows[1].position.hidden is True
    assert table.rows[-1].position.row == 6
    assert all(row.position.sheet == "订单" for row in table.rows)


def test_xlsx_multiple_sheets_require_selection():
    data = xlsx(lambda wb, ws: wb.create_sheet("other"))
    assert_error("SHEET_REQUIRED", data, spec("xlsx"))
    assert read_source(data, spec("xlsx", sheet="订单")).spec.sheet == "订单"
    assert_error("SHEET_NOT_FOUND", data, spec("xlsx", sheet="missing"))


def test_xlsx_preamble_and_only_header():
    def make(wb, ws):
        ws.delete_rows(1, 2)
        ws.append(["标题"])
        ws.append(["id", "amount"])
    table = read_source(xlsx(make), spec("xlsx", header_row=2))
    assert table.preamble_count == 1
    assert table.rows == ()
    assert_error("HEADER_ROW_MISSING", xlsx(make), spec("xlsx", header_row=3))


def test_xlsx_collects_formulas_regions_and_merged_cells():
    def make(wb, ws):
        ws["B2"] = "=1+1"
        ws["C1"] = ArrayFormula(ref="A1:C3", text="=1")
        ws["D2"] = DataTableFormula(ref="A2:D4")
        ws.merge_cells("A5:B6")
    table = read_source(xlsx(make), spec("xlsx"))
    assert table.headers[2].kind == "formula"
    assert table.rows[0].cells[1].kind == "formula"
    assert table.rows[0].cells[1].value == "=1+1"
    assert {(r.min_row, r.max_row, r.min_col, r.max_col) for r in table.formula_regions} == {(1, 3, 1, 3), (2, 2, 2, 2), (2, 4, 1, 4)}
    assert {(r.min_row, r.max_row, r.min_col, r.max_col) for r in table.merged_regions} == {(5, 6, 1, 2)}
    # Formula coverage has content semantics; styling and merges alone do not
    # extend the source range below the formula area.
    assert table.rows[-1].position.row == 4


def test_xlsx_formula_cache_is_never_used():
    data = xlsx(lambda wb, ws: setattr(ws["B2"], "value", "=1+1"))
    data = rewrite_archive(data, {"xl/worksheets/sheet1.xml": lambda source: source.replace(b"<f>1+1</f><v></v>", b"<f>1+1</f><v>999</v>")})
    table = read_source(data, spec("xlsx"))
    assert table.rows[0].cells[1].kind == "formula"
    assert table.rows[0].cells[1].value == "=1+1"


def test_xlsx_actual_bounds_limits_ignore_formatting():
    def make(wb, ws):
        ws["XFD1048576"].fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    assert len(read_source(xlsx(make), spec("xlsx", limits=Limits(max_records=1, max_columns=2))).rows) == 1
    assert_error("COLUMN_LIMIT", xlsx(lambda wb, ws: setattr(ws["C2"], "value", "extra")), spec("xlsx", limits=Limits(max_columns=2)))
    assert_error("RECORD_LIMIT", xlsx(lambda wb, ws: setattr(ws["A4"], "value", "extra")), spec("xlsx", limits=Limits(max_records=2)))
    assert_error("COLUMN_LIMIT", xlsx(lambda wb, ws: setattr(ws["C2"], "value", ArrayFormula("C2:Z4", "=1"))), spec("xlsx", limits=Limits(max_columns=5)))


def test_xlsx_rejects_external_hyperlinks_and_active_payloads():
    data = xlsx(lambda wb, ws: setattr(ws["A2"], "hyperlink", "https://example.invalid/"))
    assert_error("XLSX_EXTERNAL_LINK", data, spec("xlsx"))
    for name in ["xl/vbaProject.bin", "xl/embeddings/oleObject1.bin", "xl/connections.xml", "xl/externalLinks/externalLink1.xml"]:
        data = rewrite_archive(xlsx(), additions={name: b"payload"})
        assert_error("XLSX_ACTIVE_CONTENT", data, spec("xlsx"))


def test_xlsx_rejects_dtd_entities_even_in_unused_parts():
    data = rewrite_archive(xlsx(), additions={"docProps/unused.xml": b'<!DOCTYPE x [<!ENTITY y "payload">]><x>&y;</x>'})
    assert_error("XLSX_INVALID", data, spec("xlsx"))


def test_xlsx_zip_limits_and_invalid_archives():
    data = xlsx()
    assert_error("XLSX_ZIP_ENTRY_LIMIT", data, spec("xlsx", limits=Limits(max_zip_entries=2)))
    assert_error("XLSX_UNCOMPRESSED_LIMIT", data, spec("xlsx", limits=Limits(max_uncompressed_bytes=10)))
    assert_error("XLSX_COMPRESSION_RATIO", data, spec("xlsx", limits=Limits(max_compression_ratio=1)))
    assert_error("XLSX_INVALID", data[:-30], spec("xlsx"))
    assert_error("XLSX_CONTAINER_INVALID", rewrite_archive(data, additions={"../payload.txt": b"x"}), spec("xlsx"))
    empty = io.BytesIO()
    with zipfile.ZipFile(empty, "w") as archive:
        archive.writestr("not_a_workbook.txt", "x")
    assert_error("XLSX_CONTAINER_INVALID", empty.getvalue(), spec("xlsx"))


def test_xlsx_position_tampering_and_xml_depth_are_rejected():
    data = rewrite_archive(xlsx(), {"xl/worksheets/sheet1.xml": lambda xml: xml.replace(b'r="B2"', b'r="A2"')})
    assert_error("XLSX_POSITION_INVALID", data, spec("xlsx"))
    depth = b"<x>" * 130 + b"</x>" * 130
    data = rewrite_archive(xlsx(), additions={"docProps/unused.xml": depth})
    assert_error("XLSX_XML_DEPTH", data, spec("xlsx"))


def test_xlsx_source_ids_are_repeatable_and_sheet_sensitive():
    def make(wb, ws):
        other = wb.create_sheet("other")
        other.append(["id", "amount"])
        other.append(["00123", 12.3])
    data = xlsx(make)
    original = read_source(data, spec("xlsx", sheet="订单"))
    duplicate = read_source(data, spec("xlsx", sheet="订单"))
    other = read_source(data, spec("xlsx", sheet="other"))
    payment = read_source(data, replace(spec("xlsx", sheet="订单"), side=Side.PAYMENT))
    assert original.rows[0].source_id == duplicate.rows[0].source_id
    assert len({original.rows[0].source_id, other.rows[0].source_id, payment.rows[0].source_id}) == 3


def test_default_csv_ten_mib_file_size_exact_boundary():
    limit = 10 * 1024 * 1024
    header = b"id,amount,note\n"
    repeated = b"A,0," + b"x" * 131_000 + b"\n"
    count, remainder = divmod(limit - len(header), len(repeated))
    data = header + repeated * count + b"B,0," + b"x" * (remainder - 4)
    assert len(data) == limit
    assert len(read_source(data, spec()).rows) == count + 1
    assert_error("FILE_SIZE_LIMIT", data + b"x")


def test_default_csv_fifty_thousand_records_exact_boundary():
    data = b"id,amount\n" + b"A,0\n" * 50_000
    assert len(read_source(data, spec()).rows) == 50_000
    assert_error("RECORD_LIMIT", data + b"A,0\n")


def test_default_csv_hundred_columns_exact_boundary():
    data = b",".join([b"header"] * 100) + b"\n" + b",".join([b"value"] * 100) + b"\n"
    table = read_source(data, spec())
    assert len(table.headers) == len(table.rows[0].cells) == 100
    assert_error("COLUMN_LIMIT", b",".join([b"header"] * 101) + b"\n")


def test_csv_field_limit_cannot_be_raised_by_other_library():
    import csv
    previous = csv.field_size_limit()
    try:
        csv.field_size_limit(256 * 1024)
        prefix = b"id,amount,note\nA,0,"
        assert len(read_source(prefix + b"x" * (128 * 1024), spec()).rows) == 1
        assert_error("CSV_FIELD_LIMIT", prefix + b"x" * (128 * 1024 + 1))
    finally:
        csv.field_size_limit(previous)


def test_invalid_excel_dates_do_not_log_raw_amounts():
    import warnings
    def make(wb, ws):
        ws["B2"] = 987654321
        ws["B2"].number_format = "yyyy-mm-dd"
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        table = read_source(xlsx(make), spec("xlsx"))
    assert table.rows[0].cells[1].kind == "error"
    assert caught == []

"""M2 report acceptance: independent hand totals, safe values and complete trace.

The sample oracle is checked-in, handwritten ``samples/expected.json``.  Export
assertions are computed from that oracle or deliberately tiny source fixtures,
never by calling the writer's row-generation or formatting helpers.
"""
from __future__ import annotations

import builtins
from copy import deepcopy
import csv
from dataclasses import replace
from hashlib import sha256
import io
import json
from pathlib import Path
import socket
import tempfile
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
import pytest

from sheetcheck.engine import reconcile
from sheetcheck.errors import IntegrityError, SheetCheckError
from sheetcheck.models import FileSpec, Mapping, Rules, Side, Status
from sheetcheck.readers import read_source
from sheetcheck.report import audit_xlsx, export_xlsx
from sheetcheck.report_schema import BUSINESS_COLUMNS, ISSUE_COLUMNS, LEDGER_COLUMNS, RUN_HEADERS
from sheetcheck.validators import validate


ROOT = Path(__file__).resolve().parents[1]
EXPECTED = json.loads((ROOT / "samples" / "expected.json").read_text(encoding="utf-8"))
RULES = Rules(currency="CNY", period="synthetic", same_currency_confirmed=True,
              same_period_confirmed=True)
TABS = ["运行说明", "业务结果", "异常明细", "数据问题", "源行账本"]


def sample_run(extension="csv", *, remapped=False):
    tables = []
    for side, stem in ((Side.ORDER, "orders"), (Side.PAYMENT, "payments")):
        path = (ROOT / "tests" / "fixtures" / f"{stem}_remapped.xlsx") if remapped else ROOT / "samples" / f"{stem}.{extension}"
        spec = FileSpec(path.name, side, sheet="数据" if path.suffix == ".xlsx" else None)
        tables.append(validate(read_source(path.read_bytes(), spec),
                               Mapping(3, 2) if remapped else Mapping(1, 2), RULES))
    return reconcile(*tables, RULES)


def csv_bytes(rows):
    buffer = io.StringIO(newline="")
    csv.writer(buffer, lineterminator="\n").writerows(rows)
    return buffer.getvalue().encode("utf-8")


def csv_run(order_rows, payment_rows, *, headers=("id", "amount"),
            filenames=("orders.csv", "payments.csv"), rules=RULES):
    tables = [validate(read_source(csv_bytes([headers, *rows]), FileSpec(filename, side)),
                       Mapping(1, 2), rules)
              for side, filename, rows in zip(Side, filenames, (order_rows, payment_rows))]
    return reconcile(*tables, rules)


def workbook_bytes(workbook):
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def xlsx_run(*, header_row=1, mapping=Mapping(1, 2), unrelated_formula=False):
    """Two 0.00 rows plus an interior blank; hidden row and preamble are real."""
    tables = []
    for side in Side:
        book = Workbook()
        sheet = book.active
        sheet.title = "交易😀"
        for row in range(1, header_row):
            sheet.cell(row, 1, "前言")
        for column, value in enumerate(("id", "amount", "备注"), 1):
            sheet.cell(header_row, column, value)
        sheet.cell(header_row + 1, 1, "00001")
        sheet.cell(header_row + 1, 2, 0)
        sheet.cell(header_row + 1, 3, '=HYPERLINK("https://invalid.example/", "DO_NOT_EXPORT")'
                   if unrelated_formula else "普通备注")
        sheet.cell(header_row + 3, 1, "00002")
        sheet.cell(header_row + 3, 2, "0.00")
        sheet.row_dimensions[header_row + 1].hidden = True
        data = workbook_bytes(book)
        source = read_source(data, FileSpec(f"{side}.xlsx", side, sheet="交易😀", header_row=header_row))
        tables.append(validate(source, mapping, RULES))
    return reconcile(*tables, RULES)


def all_cells(workbook):
    return (cell for sheet in workbook for row in sheet for cell in row)


def rows_with_cells(book, name, columns):
    """Use public schema solely to locate columns, never to compute expected data."""
    sheet = book[name]
    assert tuple(cell.value for cell in sheet[1]) == tuple(column.title for column in columns)
    return [dict(zip((column.key for column in columns), cells))
            for cells in sheet.iter_rows(min_row=2)]


def row_values(row):
    return {key: cell.value for key, cell in row.items()}


def exact_money(minor):
    if minor is None:
        return None
    # Integer-only independent oracle, including negative deltas and true zero.
    absolute = abs(minor)
    return f"{'-' if minor < 0 else ''}{absolute // 100}.{absolute % 100:02d}"


def original_text(value):
    """Independent decoder for the documented SC_ESC: + JSON string protocol."""
    if isinstance(value, str) and value.startswith("SC_ESC:"):
        decoded = json.loads(value[len("SC_ESC:"):])
        assert isinstance(decoded, str)
        return decoded
    return value


def run_notes(book):
    sheet = book["运行说明"]
    assert tuple(cell.value for cell in sheet[1]) == RUN_HEADERS
    rows = list(sheet.iter_rows(min_row=2))
    assert len({row[1].value for row in rows}) == len(rows)
    return {row[1].value: row for row in rows}


def assert_passive_package(data):
    """Inspect both openpyxl interpretation and the actual serialized package."""
    book = load_workbook(io.BytesIO(data), data_only=False, keep_links=True)
    try:
        assert book.sheetnames == TABS
        assert not book._external_links
        assert book.vba_archive is None
        for cell in all_cells(book):
            assert cell.data_type != "f", cell.coordinate
            assert cell.hyperlink is None, cell.coordinate
        with ZipFile(io.BytesIO(data)) as archive:
            for name in archive.namelist():
                lower = name.lower()
                assert not any(part in lower for part in ("externallink", "vbaproject", "activex", "macrosheet"))
                if name.endswith(".rels"):
                    assert b'TargetMode="External"' not in archive.read(name)
                if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                    assert b"<f>" not in archive.read(name)
                    assert b"<hyperlink" not in archive.read(name)
    finally:
        book.close()


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_report_package_is_passive_and_public_audit_accepts_it(extension):
    result = sample_run(extension)
    data = export_xlsx(result)
    assert isinstance(data, bytes) and data.startswith(b"PK")
    assert_passive_package(data)
    assert audit_xlsx(data, result) is None


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_canonical_business_values_counts_and_money_are_hand_recomputable(extension):
    result = sample_run(extension)
    book = load_workbook(io.BytesIO(export_xlsx(result)), data_only=False)
    try:
        business = rows_with_cells(book, "业务结果", BUSINESS_COLUMNS)
        exceptions = rows_with_cells(book, "异常明细", BUSINESS_COLUMNS)
        assert len(business) == EXPECTED["manual_checks"]["business_result_rows"] == 5
        assert len(exceptions) == 3
        for row, expected in zip(business, EXPECTED["records"]):
            for key in ("status", "id", "order_minor", "payment_minor", "delta_minor", "order_row", "payment_row"):
                assert row[key].value == expected[key]
            for side in ("order", "payment", "delta"):
                number, display = row[f"{side}_minor"], row[f"{side}_amount"]
                expected_minor = expected[f"{side}_minor"]
                assert display.value == exact_money(expected_minor)
                if expected_minor is None:
                    assert number.value is None and display.value is None
                else:
                    assert type(number.value) is int
                    assert display.data_type == "s"
            assert row["id"].data_type == "s"
        assert [row_values(row) for row in exceptions] == [row_values(row) for row in business if row["status"].value != "EXACT_MATCH"]
        order_total = sum(row["order_minor"].value for row in business if row["order_minor"].value is not None)
        payment_total = sum(row["payment_minor"].value for row in business if row["payment_minor"].value is not None)
        paired_delta = sum(row["delta_minor"].value for row in business if row["delta_minor"].value is not None)
        assert (order_total, payment_total, paired_delta) == (18000, 18900, -100)
        assert payment_total - order_total == EXPECTED["manual_checks"]["valid_difference_minor"] == 900
        zero = next(row for row in business if row["id"].value == "A007")
        assert [zero[key].value for key in ("order_minor", "payment_minor", "delta_minor")] == [0, 0, 0]
        assert [zero[key].value for key in ("order_amount", "payment_amount", "delta_amount")] == ["0.00"] * 3
        only_order = next(row for row in business if row["id"].value == "A003")
        assert all(cell.value is None for key, cell in only_order.items() if key.startswith("payment_"))
        only_payment = next(row for row in business if row["id"].value == "A004")
        assert all(cell.value is None for key, cell in only_payment.items() if key.startswith("order_"))
    finally:
        book.close()


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_canonical_ledger_issues_and_business_cross_references(extension):
    result = sample_run(extension)
    book = load_workbook(io.BytesIO(export_xlsx(result)), data_only=False)
    try:
        business = [row_values(row) for row in rows_with_cells(book, "业务结果", BUSINESS_COLUMNS)]
        ledger = [row_values(row) for row in rows_with_cells(book, "源行账本", LEDGER_COLUMNS)]
        issues = [row_values(row) for row in rows_with_cells(book, "数据问题", ISSUE_COLUMNS)]
        assert len(ledger) == EXPECTED["manual_checks"]["ledger_rows"] == 14
        assert len(issues) == len(EXPECTED["issues"]) == 6
        assert len({row["source_id"] for row in ledger}) == 14
        by_id = {row["source_id"]: row for row in ledger}
        by_position = {(row["side"], row["row"]): row for row in ledger}
        for side, positions in EXPECTED["ledger_statuses"].items():
            for position, status in positions.items():
                row = by_position[(side, int(position))]
                assert row["status"] == status
                expected_filename = f"{'orders' if side == 'order' else 'payments'}.{extension}"
                assert row["filename"] == expected_filename
                assert row["sha256"] == sha256((ROOT / "samples" / expected_filename).read_bytes()).hexdigest()
                assert (row["id_column"], row["amount_column"]) == (1, 2)
                assert row["sheet"] == ("数据" if extension == "xlsx" else None)
                assert row["physical_start"] == (int(position) if extension == "csv" else None)
                assert row["physical_end"] == (int(position) if extension == "csv" else None)
        for issue, expected in zip(issues, EXPECTED["issues"]):
            assert (issue["side"], issue["row"], issue["amount_minor"]) == (expected["side"], expected["row"], expected["amount_minor"])
            assert set(expected["required_codes"]) <= set(json.loads(issue["problem_codes"]))
            assert issue["result_id"] is None
            assert issue["repair_suggestions"]
            assert {key: issue[key] for key in by_id[issue["source_id"]]} == by_id[issue["source_id"]]
        result_ids = {row["result_id"] for row in business}
        assert len(result_ids) == 5
        assert sum(row["result_id"] is not None for row in ledger) == 8
        for record in business:
            for side in ("order", "payment"):
                source_id = record[f"{side}_source_id"]
                if source_id is None:
                    continue
                source = by_id[source_id]
                assert source["result_id"] == record["result_id"]
                assert source["side"] == side
                assert source["status"] == record["status"]
                for field in ("filename", "sheet", "row", "physical_start", "physical_end", "id_column", "amount_column"):
                    assert record[f"{side}_{field}"] == source[field]
        assert all(row["result_id"] in result_ids or row["status"] == "DATA_ISSUE" for row in ledger)
    finally:
        book.close()


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_manifest_static_totals_and_complete_input_metadata(extension):
    result = sample_run(extension)
    book = load_workbook(io.BytesIO(export_xlsx(result)), data_only=False)
    try:
        notes = run_notes(book)
        assert notes["rules.currency"][2].value == "CNY"
        assert notes["rules.period"][2].value == "synthetic"
        assert notes["rules.version"][2].value == "v1"
        assert notes["rules.tolerance_minor"][2].value == 0
        assert notes["app_version"][2].value
        assert notes["report_version"][2].value
        for status, count in EXPECTED["status_counts"].items():
            cell = notes[f"status_counts.{status}"][2]
            assert cell.value == count and type(cell.value) is int
        for side, summary in EXPECTED["summaries"].items():
            for field, expected in summary.items():
                cell = notes[f"summaries.{side}.{field}"][2]
                assert cell.value == expected and type(cell.value) is int
        description = str(notes["summaries.order.issue_parseable_minor"][3].value)
        assert "不完整" in description or "无法解析" in description
        warnings = json.loads(notes["warnings"][2].value)
        assert any("缺项" in warning and "人工复核" in warning for warning in warnings)
        for side, stem in (("order", "orders"), ("payment", "payments")):
            prefix = f"inputs.{side}."
            data = (ROOT / "samples" / f"{stem}.{extension}").read_bytes()
            assert notes[prefix + "filename"][2].value == f"{stem}.{extension}"
            assert notes[prefix + "sha256"][2].value == sha256(data).hexdigest()
            assert notes[prefix + "sheet"][2].value == ("数据" if extension == "xlsx" else None)
            assert notes[prefix + "header_row"][2].value == 1
            assert notes[prefix + "mapping.id_column"][2].value == 1
            assert notes[prefix + "mapping.amount_column"][2].value == 2
            assert notes[prefix + "id_header"][2].value == "id"
            assert notes[prefix + "amount_header"][2].value == "amount"
    finally:
        book.close()


@pytest.mark.parametrize("text", [
    '=HYPERLINK("https://invalid.example/", "click")', "+SUM(1,1)", "-1+1", "@SUM(1,1)",
    "\t=1+1", "\n=1+1", "\r=1+1", "\x00=1+1", "\x1f=1+1", "\x7f=1+1", "  =1+1",
    "'literal", "SC_ESC:literal", "内嵌\x01字符", "内嵌\r回车", "多行\n中文😀", "",
])
def test_user_text_roundtrips_reversibly_without_active_cells(text):
    # Amounts may legitimately be invalid; their raw representation must still
    # survive in the issue sheet and ledger without activating spreadsheet code.
    result = csv_run([[text, text]], [])
    data = export_xlsx(result)
    assert_passive_package(data)
    book = load_workbook(io.BytesIO(data), data_only=False)
    try:
        rows = rows_with_cells(book, "源行账本", LEDGER_COLUMNS)
        assert len(rows) == 1
        for field in ("id_raw", "amount_raw"):
            cell = rows[0][field]
            assert cell.data_type == "s"
            assert original_text(cell.value) == text
        if text:
            issues = rows_with_cells(book, "数据问题", ISSUE_COLUMNS)
            assert len(issues) == 1
            assert original_text(issues[0]["id_raw"].value) == text
        else:
            assert rows[0]["status"].value == "SKIPPED_EMPTY"
        assert result.ledger[0].row.id_raw == text
        assert result.ledger[0].row.amount_raw == text
    finally:
        book.close()


def test_malicious_filename_headers_and_rule_period_are_safe_reversible_text():
    filenames = ('=HYPERLINK("x").csv', "\t@payments.csv")
    headers = ("+编号", "-金额\n中文😀")
    period = "\r=本期"
    result = csv_run([["A", "0"]], [["A", "0"]], filenames=filenames, headers=headers,
                     rules=replace(RULES, period=period))
    data = export_xlsx(result)
    assert_passive_package(data)
    book = load_workbook(io.BytesIO(data), data_only=False)
    try:
        notes = run_notes(book)
        assert original_text(notes["rules.period"][2].value) == period
        for side, filename in zip(Side, filenames):
            for field, expected in (("filename", filename), ("id_header", headers[0]), ("amount_header", headers[1])):
                cell = notes[f"inputs.{side}.{field}"][2]
                assert original_text(cell.value) == expected
                assert cell.data_type == "s"
        for row in rows_with_cells(book, "源行账本", LEDGER_COLUMNS):
            assert original_text(row["filename"].value) == filenames[0 if row["side"].value == "order" else 1]
    finally:
        book.close()


def test_unicode_128_character_identifier_and_multiline_csv_positions_survive():
    identifier = "中文😀-" + "甲" * 124
    assert len(identifier) == 128
    raw = "不是\n金额😀"
    result = csv_run([[identifier, "0001.20"], ["问题😀", raw], ["", ""], ["00042", "0"]],
                     [[identifier, "1.2"], ["00042", "0.00"]])
    book = load_workbook(io.BytesIO(export_xlsx(result)), data_only=False)
    try:
        business = rows_with_cells(book, "业务结果", BUSINESS_COLUMNS)
        assert {row["id"].value for row in business} == {identifier, "00042"}
        long_row = next(row for row in business if row["id"].value == identifier)
        assert long_row["id"].data_type == "s"
        assert long_row["order_minor"].value == long_row["payment_minor"].value == 120
        ledger = rows_with_cells(book, "源行账本", LEDGER_COLUMNS)
        assert len(ledger) == 6
        orders = [row for row in ledger if row["side"].value == "order"]
        assert [(row["row"].value, row["physical_start"].value, row["physical_end"].value)
                for row in orders] == [(2, 2, 2), (3, 3, 4), (4, 5, 5), (5, 6, 6)]
        assert original_text(orders[1]["amount_raw"].value) == raw
        assert orders[2]["status"].value == "SKIPPED_EMPTY"
        assert orders[2]["result_id"].value is None
        assert orders[2]["amount_minor"].value is None
        assert original_text(orders[0]["amount_raw"].value) == "0001.20"
        assert orders[0]["amount_raw"].data_type == "s"
    finally:
        book.close()


def test_xlsx_preamble_hidden_rows_and_interior_blanks_remain_traceable():
    result = xlsx_run(header_row=3)
    book = load_workbook(io.BytesIO(export_xlsx(result)), data_only=False)
    try:
        ledger = rows_with_cells(book, "源行账本", LEDGER_COLUMNS)
        assert len(ledger) == 6
        for side in Side:
            rows = [row for row in ledger if row["side"].value == side]
            assert [row["row"].value for row in rows] == [4, 5, 6]
            assert [row["status"].value for row in rows] == ["EXACT_MATCH", "SKIPPED_EMPTY", "EXACT_MATCH"]
            assert [row["hidden"].value for row in rows] == ["是", "否", "否"]
            assert all(row["sheet"].value == "交易😀" for row in rows)
            assert all(row["physical_start"].value is None and row["physical_end"].value is None for row in rows)
        notes = run_notes(book)
        assert notes["inputs.order.preamble_count"][2].value == 2
        assert notes["inputs.order.hidden_row_count"][2].value == 1
    finally:
        book.close()


def test_column_remapping_exports_real_column_numbers_and_original_row_positions():
    result = sample_run("xlsx", remapped=True)
    book = load_workbook(io.BytesIO(export_xlsx(result)), data_only=False)
    try:
        ledger = rows_with_cells(book, "源行账本", LEDGER_COLUMNS)
        assert len(ledger) == 14
        assert all((row["id_column"].value, row["amount_column"].value) == (3, 2) for row in ledger)
        # Open the actual remapped fixture independently of engine metadata.
        for side, stem in (("order", "orders"), ("payment", "payments")):
            source = load_workbook(ROOT / "tests" / "fixtures" / f"{stem}_remapped.xlsx", data_only=False)
            try:
                for row in ledger:
                    if row["side"].value == side:
                        actual_raw = source["数据"].cell(row["row"].value, 3).value
                        assert original_text(row["id_raw"].value) == actual_raw
            finally:
                source.close()
    finally:
        book.close()


def test_multiple_problem_codes_share_one_source_row_and_no_business_result():
    result = csv_run([["\tA", "@bad"]], [["\tA", "1.00"]])
    book = load_workbook(io.BytesIO(export_xlsx(result)), data_only=False)
    try:
        assert rows_with_cells(book, "业务结果", BUSINESS_COLUMNS) == []
        assert rows_with_cells(book, "异常明细", BUSINESS_COLUMNS) == []
        issues = rows_with_cells(book, "数据问题", ISSUE_COLUMNS)
        ledger = rows_with_cells(book, "源行账本", LEDGER_COLUMNS)
        assert len(issues) == len(ledger) == 2
        assert len({row["source_id"].value for row in issues}) == 2
        expected_codes = {
            "order": {"ID_WHITESPACE", "ID_CONTROL_CHARACTER", "AMOUNT_FORMAT", "BLOCKED_ID"},
            "payment": {"ID_WHITESPACE", "ID_CONTROL_CHARACTER", "BLOCKED_ID"},
        }
        for row in issues:
            assert set(json.loads(row["problem_codes"].value)) == expected_codes[row["side"].value]
            assert row["result_id"].value is None
    finally:
        book.close()


def test_empty_inputs_have_headers_and_zero_static_summaries():
    result = csv_run([], [])
    data = export_xlsx(result)
    audit_xlsx(data, result)
    book = load_workbook(io.BytesIO(data), data_only=False)
    try:
        for name, columns in (("业务结果", BUSINESS_COLUMNS), ("异常明细", BUSINESS_COLUMNS),
                              ("数据问题", ISSUE_COLUMNS), ("源行账本", LEDGER_COLUMNS)):
            assert rows_with_cells(book, name, columns) == []
        notes = run_notes(book)
        assert notes["summaries.order.input_count"][2].value == 0
        assert notes["summaries.payment.input_count"][2].value == 0
    finally:
        book.close()


def test_maximum_single_amount_negative_delta_and_text_displays_are_exact():
    result = csv_run([["A", "999999999.99"]], [["A", "0"]])
    book = load_workbook(io.BytesIO(export_xlsx(result)), data_only=False)
    try:
        row, = rows_with_cells(book, "业务结果", BUSINESS_COLUMNS)
        for field, expected in (("order_minor", 99_999_999_999), ("payment_minor", 0), ("delta_minor", -99_999_999_999)):
            assert row[field].value == expected and type(row[field].value) is int
        assert row["order_amount"].value == "999999999.99"
        assert row["payment_amount"].value == "0.00"
        assert row["delta_amount"].value == "-999999999.99"
        assert all(row[field].data_type == "s" for field in ("order_amount", "payment_amount", "delta_amount"))
    finally:
        book.close()


def test_16_digit_summary_uses_exact_text_instead_of_excel_numeric_rounding():
    # 10,001 identical IDs are deliberately quarantined.  The hand total is
    # larger than Excel's 15-digit numeric precision even though every source
    # amount is individually within the supported maximum.
    result = csv_run([["DUP", "999999999.99"]] * 10_001, [])
    expected = 99_999_999_999 * 10_001
    assert expected == 1_000_099_999_989_999
    data = export_xlsx(result)
    book = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    try:
        notes = run_notes(book)
        amount = notes["summaries.order.issue_parseable_minor"][2]
        assert amount.data_type == "s" and amount.value == "1000099999989999"
        assert int(amount.value) == expected
        assert "15" in notes["summaries.order.issue_parseable_minor"][3].value
        display = notes["summaries.order.issue_parseable_minor.display"][2]
        assert display.data_type == "s" and display.value == "10000999999899.99"
        assert notes["summaries.order.valid_total_minor"][2].value == 0
        assert notes["summaries.order.input_count"][2].value == 10_001
        assert notes["summaries.order.data_issue_count"][2].value == 10_001
    finally:
        book.close()


def test_export_does_not_modify_original_raw_values_or_manifest():
    result = csv_run([["=HYPERLINK(\"https://invalid.example\")", "@bad"], ["'literal", "0"]],
                     [["'literal", "0"]],
                     filenames=("=订单.csv", "@收款.csv"))
    before = deepcopy(result)
    data = export_xlsx(result)
    assert result == before
    audit_xlsx(data, result)
    assert result == before


def test_export_and_readback_need_no_disk_or_network(monkeypatch):
    result = sample_run()
    # Preloading modules and sources is outside the monitored business flow.
    # Even an in-memory final ZIP must not stage worksheet XML in a temp file.
    def forbidden(*args, **kwargs):
        raise AssertionError("Report attempted filesystem/network access")
    with monkeypatch.context() as scope:
        scope.setattr(builtins, "open", forbidden)
        scope.setattr(io, "open", forbidden)
        scope.setattr(Path, "open", forbidden)
        scope.setattr(tempfile, "NamedTemporaryFile", forbidden)
        scope.setattr(tempfile, "TemporaryFile", forbidden)
        scope.setattr(tempfile, "mkstemp", forbidden)
        scope.setattr(socket, "socket", forbidden)
        scope.setattr(socket, "create_connection", forbidden)
        data = export_xlsx(result)
        audit_xlsx(data, result)
    assert isinstance(data, bytes)


@pytest.mark.parametrize("corruption", ["missing_ledger", "duplicate_ledger", "wrong_amount", "wrong_count", "missing_warning"])
def test_invalid_result_cannot_produce_a_successful_report(corruption):
    result = sample_run()
    if corruption == "missing_ledger":
        result = replace(result, ledger=result.ledger[1:])
    elif corruption == "duplicate_ledger":
        result = replace(result, ledger=result.ledger + (result.ledger[0],))
    elif corruption == "wrong_amount":
        result = replace(result, records=(replace(result.records[0], order_minor=9999), *result.records[1:]))
    elif corruption == "wrong_count":
        counts = {**result.manifest.status_counts, "DATA_ISSUE": 0}
        result = replace(result, manifest=replace(result.manifest, status_counts=counts))
    else:
        result = replace(result, manifest=replace(result.manifest, warnings=()))
    with pytest.raises(IntegrityError):
        export_xlsx(result)


def test_unrelated_source_formula_is_never_exported():
    result = xlsx_run(unrelated_formula=True)
    data = export_xlsx(result)
    assert_passive_package(data)
    with ZipFile(io.BytesIO(data)) as archive:
        content = b"".join(archive.read(name) for name in archive.namelist() if name.endswith(".xml"))
        assert b"DO_NOT_EXPORT" not in content


def test_overlong_raw_cell_fails_closed_without_silent_excel_truncation():
    raw = "不" * 32768
    result = csv_run([["A", raw]], [])
    with pytest.raises(SheetCheckError):
        export_xlsx(result)
    assert result.ledger[0].row.amount_raw == raw

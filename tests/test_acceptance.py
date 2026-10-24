"""The oracle is hand-written expected.json, independent of implementation output."""
from __future__ import annotations

from dataclasses import asdict
import io
import csv
import json
from pathlib import Path
import random

import pytest
from openpyxl import load_workbook

from sheetcheck.audit import assert_invariants
from sheetcheck.engine import reconcile
from sheetcheck.models import FileSpec, Mapping, Rules, Side, Status
from sheetcheck.readers import read_source
from sheetcheck.validators import validate

ROOT = Path(__file__).resolve().parents[1]
EXPECTED = json.loads((ROOT / "samples" / "expected.json").read_text(encoding="utf-8"))
RULES = Rules(currency="CNY", period="synthetic", same_currency_confirmed=True, same_period_confirmed=True)


def sample_run(extension="csv", *, remapped=False):
    tables = []
    for side, stem in ((Side.ORDER, "orders"), (Side.PAYMENT, "payments")):
        path = (ROOT / "tests" / "fixtures" / f"{stem}_remapped.xlsx") if remapped else ROOT / "samples" / f"{stem}.{extension}"
        table = read_source(path.read_bytes(), FileSpec(path.name, side, sheet="数据" if path.suffix == ".xlsx" else None))
        tables.append(validate(table, Mapping(3, 2) if remapped else Mapping(1, 2), RULES))
    return reconcile(*tables, RULES)


def semantic_records(result):
    return [dict(status=str(row.status), id=row.id, order_minor=row.order_minor, payment_minor=row.payment_minor, delta_minor=row.delta_minor) for row in result.records]


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_blueprint_hand_computed_business_results(extension):
    result = sample_run(extension)
    assert_invariants(result)
    locations = {entry.source_id: entry.row.source.position.row for entry in result.ledger}
    actual = [dict(status=str(row.status), id=row.id,
                   order_row=locations.get(row.order_source_id), payment_row=locations.get(row.payment_source_id),
                   order_minor=row.order_minor, payment_minor=row.payment_minor, delta_minor=row.delta_minor)
              for row in result.records]
    assert actual == EXPECTED["records"]
    assert list(result.blocked_ids) == EXPECTED["blocked_ids"]
    assert result.manifest.unidentifiable_id_risk is True
    assert result.manifest.status_counts == EXPECTED["status_counts"]
    assert {side: asdict(value) for side, value in result.manifest.summaries.items()} == EXPECTED["summaries"]


@pytest.mark.parametrize("extension", ["csv", "xlsx"])
def test_blueprint_every_source_row_and_problem_is_traceable(extension):
    result = sample_run(extension)
    by_position = {(str(entry.row.source.side), entry.row.source.position.row): entry for entry in result.ledger}
    assert len(by_position) == len(result.ledger) == EXPECTED["manual_checks"]["ledger_rows"]
    assert len({entry.source_id for entry in result.ledger}) == len(result.ledger)
    for side, statuses in EXPECTED["ledger_statuses"].items():
        for row, status in statuses.items():
            entry = by_position[(side, int(row))]
            assert entry.status == status
            if extension == "csv":
                assert entry.row.source.position.physical_start == int(row)
                assert entry.row.source.position.physical_end == int(row)
            else:
                assert entry.row.source.position.sheet == "数据"
    issues = [entry for entry in result.ledger if entry.status == Status.DATA_ISSUE]
    assert len(issues) == len(EXPECTED["issues"])
    for expected in EXPECTED["issues"]:
        entry = by_position[(expected["side"], expected["row"])]
        assert (entry.row.id_raw or "") == expected["id_raw"]
        assert entry.row.amount_minor == expected["amount_minor"]
        assert set(expected["required_codes"]) <= set(entry.problem_codes)
        assert entry.result_id is None
    assert sum(entry.result_id is not None for entry in result.ledger) == 8


def test_manual_total_reconciliation_and_semantic_equivalence():
    csv_result = sample_run()
    xlsx_result = sample_run("xlsx")
    remapped = sample_run("xlsx", remapped=True)
    assert semantic_records(csv_result) == semantic_records(xlsx_result) == semantic_records(remapped)
    order = csv_result.manifest.summaries["order"]
    payment = csv_result.manifest.summaries["payment"]
    paired_delta = sum(record.delta_minor for record in csv_result.records if record.delta_minor is not None)
    only_payment = sum(record.payment_minor for record in csv_result.records if record.status == Status.ONLY_PAYMENT)
    only_order = sum(record.order_minor for record in csv_result.records if record.status == Status.ONLY_ORDER)
    assert payment.valid_total_minor - order.valid_total_minor == paired_delta + only_payment - only_order == 900
    assert paired_delta == -100 and only_payment == 4000 and only_order == 3000
    assert csv_result.manifest.fingerprint != xlsx_result.manifest.fingerprint
    assert xlsx_result.manifest.fingerprint != remapped.manifest.fingerprint
    original_position = {(entry.row.source.side, entry.row.normalized_id): entry.row.source.position.row for entry in xlsx_result.ledger}
    assert any(original_position.get((entry.row.source.side, entry.row.normalized_id)) != entry.row.source.position.row for entry in remapped.ledger)


def test_fixture_identifiers_are_explicit_text():
    for path in [*(ROOT / "samples").glob("*.xlsx"), *(ROOT / "tests" / "fixtures").glob("*_remapped.xlsx")]:
        workbook = load_workbook(path, data_only=False)
        sheet = workbook["数据"]
        column = 3 if "remapped" in path.stem else 1
        assert all(sheet.cell(row, column).data_type == "s" for row in range(2, sheet.max_row + 1) if sheet.cell(row, column).value is not None)
        workbook.close()


def test_repeated_inputs_have_identical_business_results_ledger_and_fingerprint():
    first, second = sample_run(), sample_run()
    assert first.records == second.records
    assert first.ledger == second.ledger
    assert first.blocked_ids == second.blocked_ids
    assert first.manifest.fingerprint == second.manifest.fingerprint
    assert first.manifest.inputs == second.manifest.inputs


@pytest.mark.parametrize("seed", list(range(12)))
def test_random_row_order_preserves_business_conclusion(seed):
    tables = []
    for side, stem in ((Side.ORDER, "orders"), (Side.PAYMENT, "payments")):
        with (ROOT / "samples" / f"{stem}.csv").open(encoding="utf-8", newline="") as handle:
            rows = list(csv.reader(handle))
        header, body = rows[0], rows[1:]
        random.Random(seed).shuffle(body)
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerows([header, *body])
        source = read_source(output.getvalue().encode(), FileSpec(f"{stem}.csv", side))
        tables.append(validate(source, Mapping(1, 2), RULES))
    result = reconcile(*tables, RULES)
    assert_invariants(result)
    assert semantic_records(result) == [{key: value for key, value in expected.items() if key not in {"order_row", "payment_row"}} for expected in EXPECTED["records"]]
    for entry in result.ledger:
        original = tables[0 if entry.row.source.side == Side.ORDER else 1].source.rows[entry.row.source.position.row - 2]
        assert original.source_id == entry.source_id
        assert original.cells[0].value == entry.row.id_raw

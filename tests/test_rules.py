"""Boundary tests assert business rules, not implementation helpers."""
from __future__ import annotations

import csv
from dataclasses import replace
from datetime import date, datetime
from decimal import Inexact, ROUND_DOWN, localcontext
import io

from openpyxl import Workbook
import pytest

from sheetcheck.audit import assert_invariants
from sheetcheck.engine import reconcile
from sheetcheck.errors import InputError, IntegrityError
from sheetcheck.models import FileSpec, Mapping, Rules, Side, Status
from sheetcheck.readers import read_source
from sheetcheck.validators import validate

RULES = Rules(currency="CNY", period="synthetic", same_currency_confirmed=True, same_period_confirmed=True)


def csv_table(rows, side=Side.ORDER):
    output = io.StringIO(newline="")
    csv.writer(output).writerows([["id", "amount"], *rows])
    return read_source(output.getvalue().encode(), FileSpec(f"{side}.csv", side))


def validated(rows, side=Side.ORDER, rules=RULES):
    return validate(csv_table(rows, side), Mapping(1, 2), rules)


def run(order_rows, payment_rows):
    result = reconcile(validated(order_rows), validated(payment_rows, Side.PAYMENT), RULES)
    assert_invariants(result)
    return result


def xlsx_table(rows, *, headers=("id", "amount"), explicit_text_ids=False, number_format=None):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "data"
    sheet.append(list(headers))
    for index, row in enumerate(rows, 2):
        sheet.append(list(row))
        if explicit_text_ids and row[0] is not None:
            sheet.cell(index, 1).data_type = "s"
        if number_format:
            sheet.cell(index, 2).number_format = number_format
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return read_source(output.getvalue(), FileSpec("test.xlsx", Side.ORDER, sheet="data"))


def test_ids_preserve_leading_zero_na_null_case_chinese_and_unicode():
    identifiers = ["00123", "123", "NA", "NULL", "A", "a", "中文-编号", "Ａ", "é", "e\u0301"]
    result = run([[key, "1"] for key in identifiers], [[key, "1.00"] for key in reversed(identifiers)])
    assert len(result.records) == len(identifiers)
    assert {row.id for row in result.records} == set(identifiers)
    assert all(row.status == Status.EXACT_MATCH for row in result.records)
    assert not result.blocked_ids


@pytest.mark.parametrize("identifier,code", [
    ("", "ID_MISSING"), ("  \t", "ID_MISSING"), ("\u3000", "ID_MISSING"),
    (" A", "ID_WHITESPACE"), ("A ", "ID_WHITESPACE"), ("\u00a0A", "ID_WHITESPACE"),
    ("A\nB", "ID_CONTROL_CHARACTER"), ("A\x00B", "ID_CONTROL_CHARACTER"),
    ("A\u200bB", "ID_CONTROL_CHARACTER"), ("A\ufeffB", "ID_CONTROL_CHARACTER"),
    ("x" * 129, "ID_TOO_LONG"),
])
def test_invalid_text_identifiers_are_not_silently_cleaned(identifier, code):
    row = validated([[identifier, "1"]]).rows[0]
    assert code in row.problem_codes
    assert row.id_raw == identifier
    assert row.normalized_id == (None if not identifier.strip() else identifier)


def test_identifier_exact_length_boundary_and_internal_space():
    table = validated([["x" * 128, "1"], ["A B", "2"]])
    assert all(not row.problem_codes for row in table.rows)


def test_invalid_identifiable_text_blocks_only_exact_same_original_id():
    result = run([[" A", "1"]], [[" A", "1"], ["A", "1"]])
    assert list(result.blocked_ids) == [" A"]
    assert [(row.id, row.status) for row in result.records] == [("A", Status.ONLY_PAYMENT)]
    assert sum(entry.status == Status.DATA_ISSUE for entry in result.ledger) == 2
    assert not result.manifest.unidentifiable_id_risk


@pytest.mark.parametrize("identifier", [123, 12345678901234567890, 1.5, True, date(2026, 1, 2), datetime(2026, 1, 2, 3), "#DIV/0!"])
def test_xlsx_nontext_identifiers_are_issues_even_if_they_look_like_ids(identifier):
    source = xlsx_table([[identifier, 1]])
    row = validate(source, Mapping(1, 2), RULES).rows[0]
    assert "ID_NOT_TEXT" in row.problem_codes
    assert row.normalized_id is None
    assert row.raw_types[0] != "text"


def test_numeric_id_display_format_cannot_restore_leading_zeros():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["id", "amount"])
    sheet.append([123, 1])
    sheet["A2"].number_format = "00000"
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    source = read_source(output.getvalue(), FileSpec("numeric-id.xlsx", Side.ORDER))
    row = validate(source, Mapping(1, 2), RULES).rows[0]
    assert "ID_NOT_TEXT" in row.problem_codes and row.normalized_id is None


def test_duplicate_including_bad_amount_isolates_both_sides_and_retains_all_problems():
    result = run([["D", "X"], ["D", "2"]], [["D", "3"]])
    assert result.records == ()
    assert result.blocked_ids == ("D",)
    assert len(result.ledger) == 3
    assert all(entry.status == Status.DATA_ISSUE and entry.result_id is None for entry in result.ledger)
    assert all("DUPLICATE_ID" in entry.problem_codes for entry in result.ledger)
    bad = next(entry for entry in result.ledger if entry.row.amount_raw == "X")
    assert "AMOUNT_FORMAT" in bad.problem_codes
    assert result.manifest.summaries["order"].issue_parseable_minor == 200
    assert result.manifest.summaries["order"].issue_unparseable_count == 1
    assert result.manifest.summaries["payment"].valid_total_minor == 0


def test_payment_side_duplicates_are_also_isolated():
    result = run([["D", "3"]], [["D", "1"], ["D", "2"]])
    assert not result.records
    assert all("DUPLICATE_ID" in entry.problem_codes for entry in result.ledger)


def test_missing_identifiers_do_not_form_a_duplicate_group():
    result = run([["", "1"], ["", "2"]], [["valid", "3"]])
    assert not result.blocked_ids
    assert result.manifest.unidentifiable_id_risk
    assert [(record.id, record.status) for record in result.records] == [("valid", Status.ONLY_PAYMENT)]
    assert all("DUPLICATE_ID" not in entry.problem_codes for entry in result.ledger)


def test_zero_is_valid_on_each_side_and_missing_side_is_not_fabricated_zero():
    result = run([["order-zero", "0.00"]], [["payment-zero", "0"]])
    assert [(row.status, row.order_minor, row.payment_minor, row.delta_minor) for row in result.records] == [
        (Status.ONLY_ORDER, 0, None, None), (Status.ONLY_PAYMENT, None, 0, None)]
    assert all(value.valid_total_minor == 0 for value in result.manifest.summaries.values())


@pytest.mark.parametrize("value,minor", [
    ("0", 0), ("0.00", 0), ("0.1", 10), ("12", 1200), ("12.3", 1230),
    ("12.30", 1230), ("0012.30", 1230), ("999999999.99", 99999999999),
])
def test_accepted_text_amounts_have_exact_integer_minor_units(value, minor):
    row = validated([["A", value]]).rows[0]
    assert row.amount_minor == minor
    assert row.problem_codes == ()
    assert type(row.amount_minor) is int


def test_amount_whitespace_normalization_is_explicit_and_raw_value_is_preserved():
    row = validated([["A", " \t12.30 \n"]]).rows[0]
    assert row.amount_minor == 1230
    assert row.amount_raw == " \t12.30 \n"
    assert "AMOUNT_TRIMMED" in row.normalization_notes
    assert not row.problem_codes


@pytest.mark.parametrize("value,code", [
    ("", "AMOUNT_MISSING"), ("   ", "AMOUNT_MISSING"),
    ("1.230", "AMOUNT_PRECISION"), ("0.001", "AMOUNT_PRECISION"),
    ("999999999.999", "AMOUNT_PRECISION"), ("1000000000", "AMOUNT_LIMIT"),
    ("1000000000.00", "AMOUNT_LIMIT"), ("-0.01", "AMOUNT_FORMAT"),
    ("-0", "AMOUNT_FORMAT"), ("+1", "AMOUNT_FORMAT"), ("(1)", "AMOUNT_FORMAT"),
    ("￥1", "AMOUNT_FORMAT"), ("$1", "AMOUNT_FORMAT"), ("1,000", "AMOUNT_FORMAT"),
    ("1%", "AMOUNT_FORMAT"), ("1e2", "AMOUNT_FORMAT"), ("NaN", "AMOUNT_FORMAT"),
    ("Infinity", "AMOUNT_FORMAT"), ("1_000", "AMOUNT_FORMAT"), (".5", "AMOUNT_FORMAT"),
    ("1.", "AMOUNT_FORMAT"), ("１.２", "AMOUNT_FORMAT"), ("١.٢", "AMOUNT_FORMAT"),
])
def test_invalid_amounts_are_never_rounded_or_coerced(value, code):
    row = validated([["A", value]]).rows[0]
    assert code in row.problem_codes
    assert row.amount_minor is None


def test_amount_parsing_ignores_callers_decimal_precision_rounding_and_traps():
    with localcontext() as context:
        context.prec = 2
        context.rounding = ROUND_DOWN
        context.traps[Inexact] = True
        rows = validated([["A", "999999999.99"], ["B", "12.30"], ["C", "1.230"]]).rows
    assert rows[0].amount_minor == 99999999999
    assert rows[1].amount_minor == 1230
    assert rows[2].amount_minor is None and "AMOUNT_PRECISION" in rows[2].problem_codes


def test_amount_parsing_ignores_callers_decimal_exponent_bounds():
    with localcontext() as context:
        context.prec = 2
        context.Emax = 2
        context.Emin = -2
        for signal in context.traps:
            context.traps[signal] = True
        rows = validated([["A", "999999999.99"], ["B", "0.01"]]).rows
    assert [row.amount_minor for row in rows] == [99999999999, 1]


def test_decimal_addition_uses_exact_integer_units():
    table = validated([["A", "0.1"], ["B", "0.2"], ["C", "0.3"]])
    assert table.rows[0].amount_minor + table.rows[1].amount_minor == table.rows[2].amount_minor == 30


@pytest.mark.parametrize("value", [True, date(2026, 1, 2), datetime(2026, 1, 2, 3), "#DIV/0!"])
def test_xlsx_boolean_date_and_error_amounts_are_not_numbers(value):
    source = xlsx_table([["A", value]], explicit_text_ids=True)
    row = validate(source, Mapping(1, 2), RULES).rows[0]
    assert "AMOUNT_TYPE" in row.problem_codes
    assert row.amount_minor is None


@pytest.mark.parametrize("value,minor", [(0, 0), (12, 1200), (12.3, 1230), (999999999.99, 99999999999)])
def test_xlsx_finite_number_amounts_use_decimal_string(value, minor):
    row = validate(xlsx_table([["A", value]], explicit_text_ids=True), Mapping(1, 2), RULES).rows[0]
    assert row.amount_minor == minor and not row.problem_codes
    assert row.raw_types == ("text", "number")


def test_xlsx_display_rounding_does_not_mask_excess_precision():
    row = validate(xlsx_table([["A", 1.234]], number_format="0.00"), Mapping(1, 2), RULES).rows[0]
    assert row.amount_minor is None and "AMOUNT_PRECISION" in row.problem_codes


@pytest.mark.parametrize("headers,rows", [
    (("=1+1", "amount"), [["A", "1"]]),
    (("id", "=1+1"), [["A", "1"]]),
    (("id", "amount"), [["=1+1", "1"]]),
    (("id", "amount"), [["A", "=1+1"]]),
])
def test_formulas_in_mapped_headers_or_values_block_whole_reconciliation(headers, rows):
    with pytest.raises(InputError):
        validate(xlsx_table(rows, headers=headers), Mapping(1, 2), RULES)


def test_unrelated_formula_is_not_evaluated_or_copied_into_business_fields():
    formula = '=HYPERLINK("https://example.invalid/private","PRIVATE_FORMULA_MARKER")'
    source = xlsx_table([["A", "1", formula]], headers=("id", "amount", "note"), explicit_text_ids=True)
    table = validate(source, Mapping(1, 2), RULES)
    assert table.rows[0].normalized_id == "A" and table.rows[0].amount_minor == 100
    result = reconcile(table, validated([["A", "1"]], Side.PAYMENT), RULES)
    assert_invariants(result)
    assert result.records[0].status == Status.EXACT_MATCH
    assert "PRIVATE_FORMULA_MARKER" not in repr(result.records)
    assert "PRIVATE_FORMULA_MARKER" not in repr(result.manifest)
    # Reader source cells may retain formula metadata for validation; no report exists in M1.


def test_wholly_empty_records_are_audited_and_do_not_create_id_risk():
    result = run([["", ""], ["", ""]], [["", ""]])
    assert not result.records and not result.blocked_ids
    assert len(result.ledger) == 3
    assert all(entry.status == Status.SKIPPED_EMPTY for entry in result.ledger)
    assert not result.manifest.unidentifiable_id_risk
    assert result.manifest.summaries["order"].skipped_empty_count == 2
    assert result.manifest.summaries["payment"].skipped_empty_count == 1


def test_whitespace_or_unmapped_content_does_not_disappear_as_empty():
    source = xlsx_table([[None, None, "note"], [" ", " ", None]], headers=("id", "amount", "note"))
    table = validate(source, Mapping(1, 2), RULES)
    assert len(table.rows) == 2
    assert all(not row.skipped_empty for row in table.rows)
    assert all("ID_MISSING" in row.problem_codes for row in table.rows)


def test_header_only_inputs_produce_empty_but_valid_ledger():
    result = run([], [])
    assert result.records == result.ledger == ()
    assert all(summary.input_count == 0 and summary.valid_total_minor == 0 for summary in result.manifest.summaries.values())
    assert not result.manifest.unidentifiable_id_risk


@pytest.mark.parametrize("contents", [b"", b"\xef\xbb\xbf", b"\n"])
def test_blank_file_cannot_supply_mappable_header(contents):
    with pytest.raises(InputError):
        validate(read_source(contents, FileSpec("empty.csv", Side.ORDER)), Mapping(1, 2), RULES)


@pytest.mark.parametrize("mapping", [Mapping(1, 1), Mapping(0, 2), Mapping(1, 3), Mapping(-1, 2)])
def test_mapping_requires_two_distinct_existing_one_based_columns(mapping):
    with pytest.raises(InputError):
        validate(csv_table([["A", "1"]]), mapping, RULES)


@pytest.mark.parametrize("rules", [
    replace(RULES, same_currency_confirmed=False), replace(RULES, same_period_confirmed=False),
    replace(RULES, currency=""), replace(RULES, period=""), replace(RULES, tolerance_minor=1),
    replace(RULES, minor_digits=3), replace(RULES, version="v2"),
    replace(RULES, max_amount_minor=100),
])
def test_v1_rules_are_confirmed_and_cannot_be_silently_relaxed(rules):
    with pytest.raises(InputError):
        validated([["A", "1"]], rules=rules)


def test_reconcile_rejects_wrong_side_order_or_mismatched_rules():
    orders = validated([["A", "1"]])
    payments = validated([["A", "1"]], Side.PAYMENT)
    with pytest.raises(InputError):
        reconcile(payments, orders, RULES)
    with pytest.raises(InputError):
        reconcile(orders, payments, replace(RULES, period="different"))


def test_integrity_checks_reject_missing_or_duplicate_ledger_source_rows():
    result = run([["A", "1"]], [["A", "1"]])
    with pytest.raises(IntegrityError):
        assert_invariants(replace(result, ledger=result.ledger[:-1]))
    with pytest.raises(IntegrityError):
        assert_invariants(replace(result, ledger=(*result.ledger, result.ledger[0])))


def test_integrity_checks_reject_corrupted_amount_and_result_link():
    result = run([["A", "1"]], [["A", "1"]])
    with pytest.raises(IntegrityError):
        assert_invariants(replace(result, records=(replace(result.records[0], order_minor=999),)))
    with pytest.raises(IntegrityError):
        assert_invariants(replace(result, ledger=(replace(result.ledger[0], result_id="missing"), result.ledger[1])))

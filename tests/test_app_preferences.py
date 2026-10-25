"""Presentation preferences must never become reconciliation inputs.

These are real Streamlit widget/state checks with fictional, in-memory files.
They do not claim browser visual, native theme, or download coverage. Tab
selection uses its session-state value because AppTest has no tab-click API.
"""
from __future__ import annotations

from hashlib import sha256
import io
import json
import re

from openpyxl import Workbook, load_workbook
import pytest

from sheetcheck.errors import InputError
from sheetcheck.models import Side
from sheetcheck.report import audit_xlsx
from sheetcheck.report_schema import SHEET_NAMES
from test_app import (confirm_rules, csv_bytes, demo_run, fresh_app, healthy,
                      set_widget, uploaded_app, widget, workflow)


CHINESE_TABS = ["准备数据", "核对结果", "使用指南"]
ENGLISH_TABS = ["Prepare data", "Results", "Guide"]
CJK = re.compile(r"[\u3400-\u9fff]")


def result_element(app):
    """Previews keep source-side order; locate results without Chinese headers."""
    assert workflow(app).result is not None
    return app.dataframe[len(workflow(app).sources)]


def visible_column_names(element):
    config = json.loads(element.proto.columns or "{}")
    return [config.get(name, {}).get("label", name) for name in element.value.columns]


def cells(table):
    return {str(value) for value in table.to_numpy().ravel()}


def take_snapshot(app):
    current = workflow(app)
    return {
        "workflow": current,
        "selections": dict(current.selections),
        "sources": dict(current.sources),
        "mappings": dict(current.mappings),
        "rules": current.rules,
        "result": current.result,
        "report": current.report_bytes,
        "report_name": current.report_name,
        "revision": current._revision,
    }


def assert_unchanged(app, before):
    current = workflow(app)
    assert current is before["workflow"]
    assert current.selections == before["selections"]
    assert current.mappings == before["mappings"]
    assert current.rules == before["rules"]
    assert current._revision == before["revision"]
    assert all(current.sources.get(side) is source for side, source in before["sources"].items())
    assert current.result is before["result"]
    assert current.report_bytes is before["report"]
    assert current.report_name == before["report_name"]
    if before["report"] is not None:
        run_id = before["result"].manifest.run_id
        assert current.report_for_run(run_id) == before["report"]


def assert_no_result(app):
    current = workflow(app)
    assert current.result is None
    assert current.report_bytes is None
    assert current.report_name is None
    assert len(app.dataframe) == len(current.sources)
    assert not app.download_button


def test_preferences_start_in_chinese_light_without_enabling_reconciliation():
    app = fresh_app()
    assert widget(app, "language").value == "zh"
    assert widget(app, "appearance").value == "light"
    assert [tab.label for tab in app.tabs] == CHINESE_TABS
    assert app.button(key="run_reconciliation").disabled
    assert_no_result(app)


def test_language_change_resends_formatted_wire_values_before_the_next_interaction():
    """Canonical state alone misses a browser retaining the old radio label."""
    app = demo_run(export=True)
    set_widget(app, "appearance", "dark")
    set_widget(app, "result_view", "金额差异")
    set_widget(app, "status_filter", ["金额差异"])
    record = next(item for item in workflow(app).result.records if item.id == "A002")
    set_widget(app, "detail_id", record.result_id)
    before = take_snapshot(app)

    for language, appearance, result_view, delimiter, guide in (
        ("en", "Dark", "Amount difference", "Comma", "Guide"),
        ("zh", "夜间", "金额差异", "逗号", "使用指南"),
    ):
        set_widget(app, "language", language)
        expected = {"appearance": appearance, "result_view": result_view,
                    "detail_id": f"A002 · {result_view}"}
        for side in ("order", "payment"):
            expected.update({f"{side}_delimiter": delimiter,
                             f"{side}_id": "1: id", f"{side}_amount": "2: amount"})
        for key, label in expected.items():
            proto = widget(app, key).proto
            assert proto.set_value, f"{key} must replace its previous locale's browser value"
            assert proto.raw_value == label
            assert proto.raw_value in proto.options
        status_proto = widget(app, "status_filter").proto
        assert status_proto.set_value
        assert list(status_proto.raw_values) == [result_view]
        assert_unchanged(app, before)

        # A later unrelated rerun must consume the new wire values successfully.
        app.session_state["workspace_tab"] = guide
        healthy(app.run())
        assert not app.error
        assert widget(app, "appearance").value == "dark"
        assert widget(app, "detail_id").value == record.result_id
        assert_unchanged(app, before)


def test_language_change_resends_unselected_mapping_and_sheet_placeholders():
    book = Workbook()
    book.active.title = "准备数据"
    book.active.append(["订单", "金额差异"])
    book.active.append(["收款", "1.00"])
    book.create_sheet("核对结果")
    content = io.BytesIO()
    book.save(content)
    book.close()
    app = fresh_app()
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    healthy(app.file_uploader[0].set_value(("准备数据.xlsx", content.getvalue(), mime)).run())
    before = take_snapshot(app)
    assert widget(app, "order_sheet").value is None
    set_widget(app, "language", "en")
    sheet = widget(app, "order_sheet")
    assert sheet.value is None and sheet.proto.set_value
    assert "准备数据" in sheet.proto.options and "核对结果" in sheet.proto.options
    assert_unchanged(app, before)

    set_widget(app, "order_sheet", "准备数据")
    before = take_snapshot(app)
    set_widget(app, "language", "zh")
    for key in ("order_id", "order_amount"):
        element = widget(app, key)
        assert element.value is None and element.proto.set_value
        assert "1: 订单" in element.proto.options
    assert widget(app, "order_sheet").proto.raw_value == "准备数据"
    assert_unchanged(app, before)
    assert_no_result(app)


def test_partial_upload_and_literal_period_survive_language_and_theme_changes():
    app = fresh_app()
    content = csv_bytes([("订单", "1.00")], headers=("准备数据", "金额差异"))
    healthy(app.file_uploader[0].set_value(("准备数据.csv", content, "text/csv")).run())
    set_widget(app, "order_id", 1)
    set_widget(app, "period", "我确认两个文件属于同一币种")
    before = take_snapshot(app)
    for key, value in (("language", "en"), ("appearance", "dark"),
                       ("language", "zh"), ("appearance", "light")):
        set_widget(app, key, value)
        assert_unchanged(app, before)
        assert widget(app, "order_id").value == 1
        assert widget(app, "order_amount").value is None
        assert widget(app, "period").value == "我确认两个文件属于同一币种"
        assert not widget(app, "confirm_currency").value
        assert not widget(app, "confirm_period").value
        assert app.button(key="run_reconciliation").disabled
        assert workflow(app).selections[Side.ORDER].data == content
        preview = app.dataframe[0].value
        assert {"1: 准备数据", "2: 金额差异"} <= set(preview.columns)
        assert "订单" in cells(preview)


def test_confirmed_inputs_keep_their_confirmation_and_do_not_run_on_translation():
    app = uploaded_app([("金额差异", "2.00")], [("金额差异", "1.00")])
    set_widget(app, "period", "中文期间保持原样")
    before = take_snapshot(app)
    set_widget(app, "language", "en")
    set_widget(app, "appearance", "dark")
    assert_unchanged(app, before)
    assert widget(app, "period").value == "中文期间保持原样"
    assert widget(app, "confirm_currency").value
    assert widget(app, "confirm_period").value
    assert not app.button(key="run_reconciliation").disabled
    assert_no_result(app)


def test_language_round_trip_preserves_the_selected_logical_workspace():
    app = demo_run(export=True)
    before = take_snapshot(app)
    for chinese, english in zip(CHINESE_TABS, ENGLISH_TABS):
        app.session_state["workspace_tab"] = chinese
        healthy(app.run())
        set_widget(app, "language", "en")
        assert [tab.label for tab in app.tabs] == ENGLISH_TABS
        assert app.session_state["workspace_tab"] == english
        set_widget(app, "language", "zh")
        assert [tab.label for tab in app.tabs] == CHINESE_TABS
        assert app.session_state["workspace_tab"] == chinese
        assert_unchanged(app, before)


def test_exported_report_filter_and_source_are_identical_across_preferences():
    app = demo_run(export=True)
    set_widget(app, "status_filter", ["金额差异"])
    set_widget(app, "query", "A002")
    record = next(item for item in workflow(app).result.records if item.id == "A002")
    set_widget(app, "detail_id", record.result_id)
    before = take_snapshot(app)
    checksum = sha256(before["report"]).hexdigest()
    for key, value in (("appearance", "dark"), ("language", "en"),
                       ("appearance", "light"), ("language", "zh")):
        set_widget(app, key, value)
        assert_unchanged(app, before)
        assert widget(app, "query").value == "A002"
        assert widget(app, "status_filter").value == ["金额差异"]
        assert widget(app, "detail_id").value == record.result_id
        assert widget(app, "page_number").value == 1
        assert len(result_element(app).value) == 1
        assert {"A002", "50.00", "49.00", "-1.00"} <= cells(result_element(app).value)
        assert sha256(workflow(app).report_bytes).hexdigest() == checksum
    audit_xlsx(workflow(app).report_bytes, workflow(app).result)
    book = load_workbook(io.BytesIO(workflow(app).report_bytes), read_only=True)
    try:
        assert tuple(book.sheetnames) == tuple(SHEET_NAMES)
    finally:
        book.close()


def test_later_page_and_source_selection_survive_language_and_theme_changes():
    rows = [(f"A{number:04d}", "1.00") for number in range(1, 206)]
    app = uploaded_app(rows, rows)
    healthy(app.button(key="run_reconciliation").click().run())
    healthy(app.button(key="prepare_report").click().run())
    set_widget(app, "status_filter", ["一致"])
    set_widget(app, "page_number", 2)
    record = workflow(app).result.records[100]
    set_widget(app, "detail_id", record.result_id)
    before = take_snapshot(app)
    set_widget(app, "language", "en")
    set_widget(app, "appearance", "dark")
    assert_unchanged(app, before)
    assert widget(app, "page_number").value == 2
    assert widget(app, "detail_id").value == record.result_id
    table = result_element(app).value
    assert len(table) == 100
    assert {f"A{number:04d}" for number in range(101, 201)} <= cells(table)
    assert "A0001" not in cells(table) and "A0201" not in cells(table)


def test_english_tables_translate_labels_but_preserve_original_chinese_values():
    identifiers = ["订单", "金额差异", "准备数据", "<script>alert(1)</script>"]
    rows = [(identifier, "1.00") for identifier in identifiers]
    app = uploaded_app(rows, rows, headers=("订单", "金额差异"))
    healthy(app.button(key="run_reconciliation").click().run())
    before = take_snapshot(app)
    set_widget(app, "language", "en")
    assert_unchanged(app, before)
    for preview in app.dataframe[:2]:
        assert {"1: 订单", "2: 金额差异"} <= set(preview.value.columns)
        assert set(identifiers) <= cells(preview.value)
    table = result_element(app)
    assert not any(CJK.search(str(name)) for name in visible_column_names(table))
    assert set(identifiers) <= cells(table.value)
    assert {record.id for record in workflow(app).result.records} == set(identifiers)


def test_chinese_worksheet_name_is_preserved_when_it_matches_ui_vocabulary():
    book = Workbook()
    book.active.title = "准备数据"
    book.active.append(["订单", "金额差异"])
    book.active.append(["收款", "1.00"])
    content = io.BytesIO()
    book.save(content)
    book.close()
    app = fresh_app()
    mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    healthy(app.file_uploader[0].set_value(("核对结果.xlsx", content.getvalue(), mime)).run())
    before = take_snapshot(app)
    set_widget(app, "language", "en")
    set_widget(app, "appearance", "dark")
    assert_unchanged(app, before)
    assert widget(app, "order_sheet").value == "准备数据"
    assert workflow(app).sources[Side.ORDER].spec.filename == "核对结果.xlsx"
    assert workflow(app).sources[Side.ORDER].spec.sheet == "准备数据"
    assert {"1: 订单", "2: 金额差异"} <= set(app.dataframe[0].value.columns)
    assert "收款" in cells(app.dataframe[0].value)


def test_reset_keeps_preferences_but_revokes_the_old_report_and_business_state():
    app = demo_run(export=True)
    old_workflow = workflow(app)
    run_id = old_workflow.result.manifest.run_id
    set_widget(app, "language", "en")
    set_widget(app, "appearance", "dark")
    app.session_state["workspace_tab"] = "Guide"
    healthy(app.run())
    healthy(app.button(key="reset").click().run())
    assert widget(app, "language").value == "en"
    assert widget(app, "appearance").value == "dark"
    assert app.session_state["workspace_tab"] == "Prepare data"
    assert widget(app, "period").value == ""
    assert not widget(app, "confirm_currency").value
    assert all(selection is None for selection in workflow(app).selections.values())
    assert_no_result(app)
    with pytest.raises(InputError, match="REPORT_STALE"):
        old_workflow.report_for_run(run_id)


def test_loading_sample_keeps_preferences_and_requires_fresh_confirmations():
    app = demo_run(export=True)
    old_workflow = workflow(app)
    run_id = old_workflow.result.manifest.run_id
    set_widget(app, "language", "en")
    set_widget(app, "appearance", "dark")
    healthy(app.button(key="load_demo").click().run())
    assert widget(app, "language").value == "en"
    assert widget(app, "appearance").value == "dark"
    assert app.session_state["workspace_tab"] == "Prepare data"
    assert widget(app, "period").value in {"虚构样例期间", "Fictional sample period"}
    assert not widget(app, "confirm_currency").value
    assert not widget(app, "confirm_period").value
    assert app.button(key="run_reconciliation").disabled
    assert all(workflow(app).selections.values())
    assert_no_result(app)
    with pytest.raises(InputError, match="REPORT_STALE"):
        old_workflow.report_for_run(run_id)


def test_replacing_an_english_mode_input_still_invalidates_and_clears_mapping():
    app = demo_run(export=True)
    old_workflow = workflow(app)
    run_id = old_workflow.result.manifest.run_id
    set_widget(app, "language", "en")
    set_widget(app, "appearance", "dark")
    data = csv_bytes([("REPLACEMENT", "3.00")])
    healthy(app.file_uploader[0].set_value(("replacement.csv", data, "text/csv")).run())
    assert_no_result(app)
    assert widget(app, "language").value == "en"
    assert widget(app, "appearance").value == "dark"
    assert app.session_state["workspace_tab"] == "Prepare data"
    assert widget(app, "order_id").value is None
    assert widget(app, "order_amount").value is None
    assert not widget(app, "confirm_currency").value
    assert workflow(app).selections[Side.ORDER].data == data
    with pytest.raises(InputError, match="REPORT_STALE"):
        old_workflow.report_for_run(run_id)


def test_english_issue_table_keeps_problem_codes_and_raw_amounts():
    app = demo_run(export=True)
    set_widget(app, "result_view", "数据问题")
    set_widget(app, "query", "A006")
    before = take_snapshot(app)
    set_widget(app, "language", "en")
    assert_unchanged(app, before)
    table = result_element(app)
    assert not any(CJK.search(str(name)) for name in visible_column_names(table))
    assert len(table.value) == 2
    values = cells(table.value)
    assert {"A006", "X", "12.00", "BLOCKED_ID"} <= values
    assert any("AMOUNT_FORMAT" in value for value in values)
    # These sample rows contain no original Chinese data, so localized
    # roles, statuses, explanations and repair advice must all be English.
    assert not any(CJK.search(value) for value in values)

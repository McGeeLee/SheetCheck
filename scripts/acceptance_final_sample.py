"""Explicitly export the fictional CLI sample and independently check its oracle.

This acceptance script writes only the requested evidence directory and refuses
to reuse one. It does not modify the source files or the hand-written expected
answers. Run with the installed environment from the project root.
"""
from __future__ import annotations

import argparse
from dataclasses import asdict, replace
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
from pathlib import Path
import subprocess
import sys

from openpyxl import load_workbook

from sheetcheck import FileSpec, Mapping, Rules, Side, read_source, reconcile, validate
from sheetcheck.__main__ import result_to_dict
from sheetcheck.audit import assert_invariants
from sheetcheck.models import APP_VERSION, Status
from sheetcheck.report import audit_xlsx, report_filename
from sheetcheck.report_schema import SHEET_NAMES

ROOT = Path(__file__).resolve().parents[1]


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    if Path.cwd().resolve() != ROOT:
        raise SystemExit("Run from the SheetCheck project directory.")
    destination = args.output_dir.resolve()
    destination.mkdir(parents=True, exist_ok=False)
    protected = [ROOT / "samples" / name for name in ("orders.csv", "payments.csv", "expected.json")]
    before = {str(path.relative_to(ROOT)): digest(path) for path in protected}
    expected = json.loads(protected[2].read_text(encoding="utf-8"))
    command = [sys.executable, "-m", "sheetcheck", "--demo", "--export-dir", str(destination)]
    completed = subprocess.run(command, capture_output=True, text=True, check=False)
    (destination / "cli-stderr.txt").write_text(completed.stderr, encoding="utf-8")
    if completed.returncode != 0:
        raise SystemExit(f"CLI failed with exit code {completed.returncode}; no successful evidence issued.")
    payload = json.loads(completed.stdout)
    assert completed.stderr == f"code=EXPORTED stage=EXPORT version={APP_VERSION}\n"
    assert payload["manifest"]["app_version"] == APP_VERSION
    (destination / "cli-demo.json").write_text(completed.stdout, encoding="utf-8")

    # Re-read both original files in this independent parent process. Only
    # nondeterministic run identity/timestamps are carried over from CLI output.
    rules = Rules(currency=expected["currency"], period=expected["period"],
                  same_currency_confirmed=True, same_period_confirmed=True)
    tables = [validate(read_source(path.read_bytes(), FileSpec(path.name, side)), Mapping(1, 2), rules)
              for side, path in zip(Side, protected[:2], strict=True)]
    result = reconcile(*tables, rules)
    assert result.manifest.fingerprint == payload["manifest"]["fingerprint"]
    result = replace(result, manifest=replace(result.manifest, **{
        name: payload["manifest"][name] for name in ("run_id", "started_at", "finished_at")}))
    assert json.loads(json.dumps(result_to_dict(result))) == payload
    assert_invariants(result)

    positions = {entry.source_id: entry.row.source.position.row for entry in result.ledger}
    records = [dict(status=str(row.status), id=row.id,
                    order_row=positions.get(row.order_source_id), payment_row=positions.get(row.payment_source_id),
                    order_minor=row.order_minor, payment_minor=row.payment_minor, delta_minor=row.delta_minor)
               for row in result.records]
    assert records == expected["records"]
    assert list(result.blocked_ids) == expected["blocked_ids"]
    assert result.manifest.unidentifiable_id_risk == expected["unidentifiable_id_risk"]
    assert result.manifest.status_counts == expected["status_counts"]
    assert {str(side): asdict(value) for side, value in result.manifest.summaries.items()} == expected["summaries"]
    by_position = {(str(entry.row.source.side), entry.row.source.position.row): entry for entry in result.ledger}
    assert len(by_position) == len(result.ledger) == expected["manual_checks"]["ledger_rows"]
    assert len(set(entry.source_id for entry in result.ledger)) == len(result.ledger)
    for side, statuses in expected["ledger_statuses"].items():
        for row, status in statuses.items():
            entry = by_position[(side, int(row))]
            assert entry.status == status
            assert entry.row.source.position.physical_start == int(row)
            assert entry.row.source.position.physical_end == int(row)
    issues = [entry for entry in result.ledger if entry.status == Status.DATA_ISSUE]
    assert len(issues) == len(expected["issues"])
    for answer in expected["issues"]:
        entry = by_position[(answer["side"], answer["row"])]
        assert entry.row.id_raw == answer["id_raw"]
        assert entry.row.amount_minor == answer["amount_minor"]
        assert set(answer["required_codes"]) <= set(entry.problem_codes)
        assert entry.result_id is None
    manual = expected["manual_checks"]
    assert sum(entry.result_id is not None for entry in result.ledger) == manual["business_source_rows"]
    assert len(records) == manual["business_result_rows"]
    paired = sum(row.delta_minor for row in result.records if row.delta_minor is not None)
    only_payment = sum(row.payment_minor for row in result.records if row.status == Status.ONLY_PAYMENT)
    only_order = sum(row.order_minor for row in result.records if row.status == Status.ONLY_ORDER)
    difference = result.manifest.summaries[Side.PAYMENT].valid_total_minor - result.manifest.summaries[Side.ORDER].valid_total_minor
    assert paired == manual["paired_delta_minor"]
    assert only_payment == manual["only_payment_minor"]
    assert only_order == manual["only_order_minor"]
    assert difference == paired + only_payment - only_order == manual["valid_difference_minor"]

    reports = list(destination.glob("*.xlsx"))
    assert len(reports) == 1 and reports[0].name == report_filename(result)
    report = reports[0]
    data = report.read_bytes()
    audit_xlsx(data, result)
    workbook = load_workbook(BytesIO(data), data_only=False, keep_links=True)
    try:
        assert tuple(workbook.sheetnames) == SHEET_NAMES
        assert not workbook._external_links and workbook.vba_archive is None
        counts = {sheet.title: sheet.max_row - 1 for sheet in workbook}
        for sheet in workbook:
            assert sheet.sheet_state == "visible"
            for row in sheet:
                assert all(cell.data_type != "f" and cell.hyperlink is None for cell in row)
        assert counts["业务结果"] == len(expected["records"])
        assert counts["异常明细"] == 3
        assert counts["数据问题"] == len(expected["issues"])
        assert counts["源行账本"] == manual["ledger_rows"]
    finally:
        workbook.close()
    after = {str(path.relative_to(ROOT)): digest(path) for path in protected}
    assert before == after
    evidence = {
        "checked_at_utc": datetime.now(timezone.utc).isoformat(), "status": "PASS",
        "application_version": APP_VERSION, "command": command, "cli_exit_code": completed.returncode,
        "scope": "Actual fictional CLI export, independent source re-read, hand-written oracle, full JSON equality and XLSX audit",
        "report": str(report.relative_to(ROOT)), "report_bytes": len(data), "report_sha256": digest(report),
        "json": str((destination / "cli-demo.json").relative_to(ROOT)),
        "json_sha256": digest(destination / "cli-demo.json"), "fingerprint": result.manifest.fingerprint,
        "run_id": result.manifest.run_id, "data_rows_by_sheet": counts,
        "status_counts": result.manifest.status_counts,
        "summaries": {str(side): asdict(value) for side, value in result.manifest.summaries.items()},
        "manual_checks": manual, "protected_input_and_oracle_hashes": after,
        "source_files_unchanged": True, "no_formula_hyperlink_external_link_or_macro": True,
        "source_reconciliation_and_full_report_audit": "PASS",
    }
    (destination / "acceptance.json").write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"PASS: CLI export, independent sample oracle and readback; {len(result.ledger)} source rows.")


if __name__ == "__main__":
    main()

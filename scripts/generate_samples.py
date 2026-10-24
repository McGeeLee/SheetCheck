"""Rebuild fictional XLSX fixtures from the reviewed CSV files; never write expected.json.

Run from any directory: python scripts/generate_samples.py
This helper has no dependency on the reconciliation engine.
"""
from __future__ import annotations

import csv
from pathlib import Path
import random

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]


def write_workbook(path: Path, rows: list[list[str]], *, remapped: bool = False) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "数据"
    if remapped:
        sheet.append(["说明", "金额（CNY）", "业务编号"])
        random.Random(20261023).shuffle(rows)
        for index, row in enumerate(rows, 2):
            sheet.cell(index, 1, "虚构验收数据").data_type = "s"
            sheet.cell(index, 2, row[1]).data_type = "s"
            sheet.cell(index, 3, row[0]).data_type = "s"
            sheet.cell(index, 3).number_format = "@"
    else:
        sheet.append(["id", "amount"])
        for index, row in enumerate(rows, 2):
            sheet.cell(index, 1, row[0]).data_type = "s"
            sheet.cell(index, 1).number_format = "@"
            sheet.cell(index, 2, row[1]).data_type = "s"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    workbook.close()


def main() -> None:
    for name in ("orders", "payments"):
        with (ROOT / "samples" / f"{name}.csv").open(newline="", encoding="utf-8") as handle:
            table = list(csv.reader(handle))
        assert table[0] == ["id", "amount"]
        write_workbook(ROOT / "samples" / f"{name}.xlsx", table[1:])
        write_workbook(ROOT / "tests" / "fixtures" / f"{name}_remapped.xlsx", table[1:], remapped=True)
    print("Wrote four fictional XLSX fixtures; hand-written expected.json was not modified.")


if __name__ == "__main__":
    main()
